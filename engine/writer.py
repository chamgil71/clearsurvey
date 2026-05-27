from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from engine.config import ColumnDef, SurveyConfig
from engine.styler import build_cleaned_styles, load_style
from transforms.registry import TransformRegistry

MAX_DATA_ROW = 10000

# transforms that should use SUBTOTAL(9,…) = SUM instead of COUNTA
_NUMERIC_ROLES = {"n_jang", "jang", "o_binary", "to_binary", "norm_num", "normalize_number", "to_numeric"}

# transforms that return a dict → multiple derived output columns
# key: transform name  /  value: ordered list of column suffixes
# "" = main column (output_col as-is), others appended to output_col
_DERIVED_SUFFIXES: dict[str, list[str]] = {
    "norm_date_parts": ["", "_년", "_월", "_일"],
    "addr_split":      ["", "_시도", "_시군구", "_상세"],
}


def _resolve_val(row: pd.Series, col_def: ColumnDef) -> Any:
    if col_def.source_cols is not None:
        return [
            (row.iloc[sc - 1] if sc - 1 < len(row) else None)
            for sc in col_def.source_cols
        ]
    if col_def.source_col is not None:
        pos = col_def.source_col - 1
        return row.iloc[pos] if pos < len(row) else None
    if col_def.source_label is not None:
        return row.get(col_def.source_label)
    return None


def _apply_transform(
    col_def: ColumnDef,
    row: pd.Series,
    registry: TransformRegistry,
    extra_kwargs: dict,
) -> Any:
    raw = _resolve_val(row, col_def)
    if col_def.transform is None or col_def.transform not in registry:
        return raw if raw is not None and str(raw).strip() != "" else None

    kw = dict(extra_kwargs)
    if col_def.flag_keyword is not None:
        kw["flag_keyword"] = col_def.flag_keyword
    if col_def.backup_col is not None:
        bkup_pos = col_def.backup_col - 1
        kw["backup_val"] = row.iloc[bkup_pos] if bkup_pos < len(row) else None
    if col_def.year_col is not None:
        yr_pos = col_def.year_col - 1
        kw["year_val"] = row.iloc[yr_pos] if yr_pos < len(row) else None

    return registry.apply(col_def.transform, raw, **kw)


class CleanedSheetWriter:
    def __init__(self, cfg: SurveyConfig, registry: TransformRegistry):
        self._cfg      = cfg
        self._registry = registry

    def _expand_cols(
        self, col_defs: list[ColumnDef]
    ) -> list[tuple[ColumnDef, str]]:
        """Expand ColumnDefs with derived suffixes for multi-output transforms.

        Returns list of (ColumnDef, suffix) pairs.
        "" suffix = main column (output_col unchanged).
        Non-"" suffix is appended to output_col name.
        """
        expanded: list[tuple[ColumnDef, str]] = []
        for cd in col_defs:
            suffixes = _DERIVED_SUFFIXES.get(cd.transform or "", [""])
            for sfx in suffixes:
                expanded.append((cd, sfx))
        return expanded

    def write(self, wb: Workbook, df: pd.DataFrame) -> tuple[int, dict[str, int], dict[str, list]]:
        """Write Cleaned sheet into wb from preprocessed DataFrame.

        Returns (row_count, col_index_map, cleaned_col_vals).
        col_index_map: {output_col_name: 1-based column index} — includes derived cols.
        """
        style = load_style(None, self._cfg.style_file)
        s     = build_cleaned_styles(style)

        sheet_name = self._cfg.sheets.cleaned
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)

        base_defs = [cd for cd in self._cfg.columns if cd.transform != "exclude"]
        expanded  = self._expand_cols(base_defs)   # [(ColumnDef, suffix), ...]
        n_cols    = len(expanded)
        last_letter = get_column_letter(n_cols)

        extra_kw = dict(self._cfg.transform_kwargs)
        if self._cfg.jang_extraction:
            extra_kw["jang_cfg"] = self._cfg.jang_extraction.model_dump()

        # ── Row 1: SUBTOTAL aggregation row ─────────────────────────────────
        for ci, (cd, sfx) in enumerate(expanded, 1):
            letter = get_column_letter(ci)
            rng    = f"{letter}3:{letter}{MAX_DATA_ROW}"
            fn     = 9 if cd.transform in _NUMERIC_ROLES else 3
            c = ws.cell(1, ci, f"=SUBTOTAL({fn},{rng})")
            c.fill, c.font, c.alignment, c.border = (
                s["sub_fill"], s["sub_font"], s["sub_align"], s["sub_bdr"]
            )
        ws.row_dimensions[1].height = s["sub_height"]

        # ── Row 2: header labels ─────────────────────────────────────────────
        for ci, (cd, sfx) in enumerate(expanded, 1):
            header = cd.output_col + sfx
            c = ws.cell(2, ci, header)
            c.fill, c.font, c.alignment = s["hdr_fill"], s["hdr_font"], s["hdr_align"]
        ws.row_dimensions[2].height = s["hdr_height"]

        # ── Rows 3+: data ────────────────────────────────────────────────────
        row_out = 3
        n_total = len(df)
        _PROG   = 1000
        cleaned_col_vals: dict[str, list] = {
            cd.output_col + sfx: [] for (cd, sfx) in expanded
        }
        for ri, (_, row) in enumerate(df.iterrows()):
            if n_total >= _PROG and ri > 0 and ri % _PROG == 0:
                pct = int(ri / n_total * 100)
                print(f"  변환중: {ri:,}/{n_total:,}행 ({pct}%) ...", end="\r", flush=True)

            ci = 1
            for cd in base_defs:
                val = _apply_transform(cd, row, self._registry, extra_kw)
                suffixes = _DERIVED_SUFFIXES.get(cd.transform or "", [""])
                for sfx in suffixes:
                    if isinstance(val, dict):
                        cell_val = val.get(sfx)
                    elif sfx == "":
                        cell_val = val
                    else:
                        cell_val = None
                    col_key = cd.output_col + sfx
                    cleaned_col_vals[col_key].append(cell_val)
                    cell = ws.cell(row_out, ci, cell_val)
                    cell.font      = s["dat_font"]
                    cell.alignment = s["dat_align"]
                    if cd.number_format and sfx == "":
                        cell.number_format = cd.number_format
                    if cd.align and sfx == "":
                        from openpyxl.styles import Alignment as _Align
                        cell.alignment = _Align(horizontal=cd.align, vertical="center")
                    ci += 1
            row_out += 1

        if n_total >= _PROG:
            print(f"  변환중: {n_total:,}/{n_total:,}행 (100%) 완료    ")

        last_row = max(row_out - 1, 2)

        # ── Excel Table ──────────────────────────────────────────────────────
        table_ref = f"A2:{last_letter}{last_row}"
        tab = Table(displayName="Survey", ref=table_ref)
        tab.tableStyleInfo = s["tbl_style"]
        ws.add_table(tab)

        # ── column widths ────────────────────────────────────────────────────
        for ci, (cd, sfx) in enumerate(expanded, 1):
            ws.column_dimensions[get_column_letter(ci)].width = cd.width or 14

        ws.freeze_panes = s["freeze_pane"]

        col_index_map = {
            cd.output_col + sfx: ci
            for ci, (cd, sfx) in enumerate(expanded, 1)
        }
        return row_out - 3, col_index_map, cleaned_col_vals
