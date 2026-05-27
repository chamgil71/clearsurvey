from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from engine.config import SummarySection, SurveyConfig
from engine.styler import build_summary_styles, load_style

MAX_DATA_ROW = 10000


def _data_ref(cleaned: str, letter: str) -> str:
    return f"'{cleaned}'!{letter}3:{letter}{MAX_DATA_ROW}"


# ---------------------------------------------------------------------------
# Layout helpers
# ---------------------------------------------------------------------------

def _section_height(
    sec: SummarySection,
    df: pd.DataFrame,
    col_index_map: dict[str, int],
    cleaned_col_vals: dict[str, list] | None = None,
) -> int:
    """Excel rows a section occupies: title(1) + header(1) + data rows."""
    if sec.type == "unique_count":
        col_ref = sec.col_ref or ""
        if cleaned_col_vals and col_ref in cleaned_col_vals:
            vals = [str(v).strip() for v in cleaned_col_vals[col_ref] if v is not None]
            n = len({v for v in vals if v})
        else:
            n = 0
        return 2 + n
    if sec.type == "totals":
        return 2 + len(sec.items)
    if sec.type in ("binary_sum", "gpu_demand"):
        return 2 + len(sec.columns)
    if sec.type == "countif_contains":
        return 2 + len(sec.keywords)
    return 5


def _resolve_positions(
    cfg: SurveyConfig,
    df: pd.DataFrame,
    col_index_map: dict[str, int],
    cleaned_col_vals: dict[str, list] | None = None,
) -> list[SummarySection]:
    """Assign start_row / start_col to every section via the layout grid.

    - If a section already has both start_row and start_col set → used as-is.
    - If no layout is configured → sections must have explicit positions.
    """
    layout = cfg.summary.layout
    if layout is None:
        return cfg.summary.sections

    # row cursor per layout column
    cursors: dict[int, int] = {lc: layout.start_row for lc in range(1, layout.cols + 1)}

    def _start_col(layout_col: int) -> int:
        return 1 + (layout_col - 1) * (layout.col_span + layout.gap_cols)

    resolved: list[SummarySection] = []
    for sec in cfg.summary.sections:
        if sec.start_row is not None and sec.start_col is not None:
            # manual override — keep as-is but advance cursor if needed
            resolved.append(sec)
            lc = sec.layout_col
            h  = _section_height(sec, df, col_index_map)
            if cursors.get(lc, 0) <= sec.start_row:
                cursors[lc] = sec.start_row + h + layout.gap_rows
            continue

        lc     = sec.layout_col
        row    = cursors.get(lc, layout.start_row)
        col    = _start_col(lc)
        height = _section_height(sec, df, col_index_map, cleaned_col_vals)
        cursors[lc] = row + height + layout.gap_rows

        resolved.append(sec.model_copy(update={"start_row": row, "start_col": col}))

    return resolved


# ---------------------------------------------------------------------------
# Writer
# ---------------------------------------------------------------------------

class SummarySheetWriter:
    def __init__(self, cfg: SurveyConfig):
        self._cfg = cfg

    def write(
        self,
        wb: Workbook,
        df: pd.DataFrame,
        col_index_map: dict[str, int],
        cleaned_col_vals: dict[str, list] | None = None,
    ) -> None:
        style = load_style(None, self._cfg.style_file)
        s     = build_summary_styles(style)

        cleaned_name = self._cfg.sheets.cleaned
        sheet_name   = self._cfg.summary.sheet_name

        # resolve positions
        sections = _resolve_positions(self._cfg, df, col_index_map, cleaned_col_vals)

        # auto-compute total_cell from the "totals" section when layout is active
        totals_sec = next((sec for sec in sections if sec.type == "totals"), None)
        if totals_sec and self._cfg.summary.layout and totals_sec.start_row is not None:
            tc_col     = get_column_letter((totals_sec.start_col or 1) + 1)
            tc_row     = (totals_sec.start_row) + 2  # title+header+first data row
            total_cell = f"${tc_col}${tc_row}"
        else:
            total_cell = self._cfg.summary.total_cell or "$B$3"

        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name, 0)

        sec_fill = s["sec_fill"]
        sec_font = s["sec_font"]
        hdr_fill = s["hdr_fill"]
        hdr_font = s["hdr_font"]
        dat_font = s["dat_font"]
        ctr      = s["ctr"]
        lft      = s["lft"]
        bdr      = s["bdr"]

        def _sec_title(row: int, col: int, title: str, span: int = 3) -> None:
            c = ws.cell(row, col, title)
            c.fill, c.font, c.alignment = sec_fill, sec_font, ctr
            if span > 1:
                ws.merge_cells(start_row=row, start_column=col,
                               end_row=row, end_column=col + span - 1)

        def _hdr(row: int, col: int, labels: list[str]) -> None:
            for i, lbl in enumerate(labels):
                c = ws.cell(row, col + i, lbl)
                c.fill, c.font, c.alignment, c.border = hdr_fill, hdr_font, ctr, bdr

        def _dval(row: int, col: int, val: object) -> None:
            c = ws.cell(row, col, val)
            c.font, c.alignment, c.border = dat_font, lft, bdr

        def _fval(row: int, col: int, formula: str) -> None:
            c = ws.cell(row, col, formula)
            c.font, c.alignment, c.border = dat_font, ctr, bdr

        def _col_letter(col_ref: str) -> str:
            idx = col_index_map.get(col_ref, 1)
            return get_column_letter(idx)

        def _unique_vals_from_df(col_ref: str) -> list[str]:
            if cleaned_col_vals and col_ref in cleaned_col_vals:
                vals = [str(v).strip() for v in cleaned_col_vals[col_ref] if v is not None]
                return sorted({v for v in vals if v})
            # fallback: use raw df (legacy path, may be inaccurate for transforms)
            idx = col_index_map.get(col_ref, 1) - 1
            if idx < len(df.columns):
                series = df.iloc[:, idx].dropna().astype(str).str.strip()
            else:
                series = pd.Series(dtype=str)
            return sorted(v for v in series.unique() if v)

        for sec in sections:
            srow = sec.start_row
            scol = sec.start_col

            # ── totals ────────────────────────────────────────────────────────
            if sec.type == "totals":
                _sec_title(srow, scol, sec.title, span=3)
                _hdr(srow + 1, scol, ["항목", "값", "비율(%)"])
                for i, item in enumerate(sec.items):
                    row = srow + 2 + i
                    ltr = _col_letter(item.col_ref)
                    rng = _data_ref(cleaned_name, ltr)
                    _dval(row, scol, item.label)
                    if item.type == "count_all":
                        _fval(row, scol + 1, f"=COUNTA({rng})")
                        _fval(row, scol + 2, "")
                    elif item.type == "countif_exact":
                        v = (item.value or "").replace('"', '""')
                        _fval(row, scol + 1, f'=COUNTIF({rng},"{v}")')
                        cl = get_column_letter(scol + 1)
                        _fval(row, scol + 2,
                              f'=IFERROR(ROUND({cl}{row}/{total_cell}*100,1),"")')

            # ── unique_count ──────────────────────────────────────────────────
            elif sec.type == "unique_count":
                col_ref = sec.col_ref or ""
                ltr     = _col_letter(col_ref)
                rng     = _data_ref(cleaned_name, ltr)
                vals    = _unique_vals_from_df(col_ref)
                if sec.sort:
                    vals = sorted(vals)
                _sec_title(srow, scol, sec.title, span=3)
                _hdr(srow + 1, scol, ["항목", "응답수", "비율(%)"])
                for i, val in enumerate(vals):
                    row = srow + 2 + i
                    safe = val.replace('"', '""')
                    _dval(row, scol, val)
                    _fval(row, scol + 1, f'=COUNTIF({rng},"{safe}")')
                    cl = get_column_letter(scol + 1)
                    _fval(row, scol + 2,
                          f'=IFERROR(ROUND({cl}{row}/{total_cell}*100,1),"")')

            # ── binary_sum ────────────────────────────────────────────────────
            elif sec.type == "binary_sum":
                _sec_title(srow, scol, sec.title, span=3)
                _hdr(srow + 1, scol, ["항목", "응답수(중복포함)", "비율(%)"])
                for i, item in enumerate(sec.columns):
                    row = srow + 2 + i
                    ltr = _col_letter(item.col_ref)
                    rng = _data_ref(cleaned_name, ltr)
                    _dval(row, scol, item.label or item.col_ref)
                    _fval(row, scol + 1, f"=IFERROR(SUM({rng}),0)")
                    cl = get_column_letter(scol + 1)
                    _fval(row, scol + 2,
                          f'=IFERROR(ROUND({cl}{row}/{total_cell}*100,1),"")')

            # ── gpu_demand ────────────────────────────────────────────────────
            elif sec.type == "gpu_demand":
                _sec_title(srow, scol, sec.title, span=4)
                _hdr(srow + 1, scol, ["GPU 모델", "총 장수(합)", "요청건수", "건당 평균"])
                for i, item in enumerate(sec.columns):
                    row = srow + 2 + i
                    ltr = _col_letter(item.col_ref)
                    rng = _data_ref(cleaned_name, ltr)
                    c2  = get_column_letter(scol + 1)
                    c3  = get_column_letter(scol + 2)
                    _dval(row, scol, item.label or item.col_ref)
                    _fval(row, scol + 1, f'=SUMIF({rng},">0")')
                    _fval(row, scol + 2, f'=COUNTIF({rng},">0")')
                    _fval(row, scol + 3,
                          f'=IFERROR(ROUND({c2}{row}/{c3}{row},1),"")')

            # ── countif_contains ──────────────────────────────────────────────
            elif sec.type == "countif_contains":
                col_ref = sec.col_ref or ""
                ltr     = _col_letter(col_ref)
                rng     = _data_ref(cleaned_name, ltr)
                _sec_title(srow, scol, sec.title, span=3)
                _hdr(srow + 1, scol, ["항목", "응답수(중복포함)", "비율(%)"])
                for i, kw_item in enumerate(sec.keywords):
                    row = srow + 2 + i
                    _dval(row, scol, kw_item.label)
                    _fval(row, scol + 1, f'=COUNTIF({rng},"*{kw_item.keyword}*")')
                    cl = get_column_letter(scol + 1)
                    _fval(row, scol + 2,
                          f'=IFERROR(ROUND({cl}{row}/{total_cell}*100,1),"")')

        for col_letter, w in s["col_widths"].items():
            ws.column_dimensions[col_letter].width = w

        ws.freeze_panes = s["freeze_pane"]
