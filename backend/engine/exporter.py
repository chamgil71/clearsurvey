"""Export cleaned.xlsx → survey_data.json for web dashboard.

구조 (2026-07-17 분리):
  _read_cleaned_sheet()  xlsx 를 아는 유일한 함수. Cleaned 시트 → (headers, rows)
  build_data_json()      rows → 대시보드 JSON. **xlsx 를 열지 않는다.**
  write_data_json()      JSON → 파일 + manifest 갱신
  export_to_json()       위 셋을 잇는 얇은 껍데기 (기존 호출부 호환)

`build_data_json()` 이 분리된 이유: 대시보드 셀 편집이 저장될 때 xlsx 전체를 다시 쓰지 않고
data.json 만 갱신할 수 있어야 한다. 그때 타입 감지·aggregates 를 재계산해야 하는데, 그 로직이
xlsx 읽기와 한 함수에 묶여 있으면 편집 경로가 xlsx 를 열 수밖에 없다.
두 경로가 같은 함수를 쓰므로 산출물이 어긋날 수 없다.
상세: docs/plan/pending/dashboard_edit_plan.md §4.2
"""
from __future__ import annotations

import json
import math
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl

from engine.config import ROW_ID_COL, SurveyConfig


# ---------------------------------------------------------------------------
# Type detection
# ---------------------------------------------------------------------------

def _detect_type(values: list[Any]) -> str:
    if not values:
        return "text"
    non_null = [v for v in values if v is not None]
    if not non_null:
        return "text"

    numeric_count = sum(1 for v in non_null if isinstance(v, (int, float)))
    if numeric_count / len(non_null) >= 0.8:
        return "numeric"

    unique_strs = {str(v).strip() for v in non_null if str(v).strip()}
    ratio = len(unique_strs) / len(non_null)
    if ratio <= 0.5 and len(unique_strs) <= 40:
        return "category"

    return "text"


# ---------------------------------------------------------------------------
# Main export function
# ---------------------------------------------------------------------------

def _build_default_dashboard(columns: list[dict]) -> dict:
    """컬럼 메타에서 React DashboardConfig 호환 기본 설정을 생성합니다.

    React 타입 (web/src/types/dashboard.ts) 과 동일한 구조를 반환합니다:
      version: number
      kpi:    [{label, type: "total_rows"|"sum"|"count_value", col?, value?}]
      charts: [{col, type: "donut"|"bar"|"hbar"|"multibar", title?}]
      list:   {visible_cols: [...], filter_cols: [...]}
    """
    cat_cols = [c for c in columns if c["type"] == "category"]
    num_cols = [c for c in columns if c["type"] == "numeric"]

    # ── KPI ────────────────────────────────────────────────────────────────
    kpi: list[dict] = [{"label": "총 응답수", "type": "total_rows"}]
    for col in num_cols[:1]:
        kpi.append({"label": col["label"] + " 합계", "type": "sum", "col": col["key"]})

    # ── 차트 ───────────────────────────────────────────────────────────────
    charts: list[dict] = []
    for i, col in enumerate(cat_cols[:5]):
        uc = col.get("unique_count", 0)
        chart_type = "donut" if uc <= 5 else ("bar" if uc <= 15 else "hbar")
        charts.append({"col": col["key"], "type": chart_type, "title": col["label"]})

    # O_ 접두사 이진 컬럼 → multibar 하나로 묶기
    binary_cols = [c for c in num_cols if c["key"].startswith("O_")]
    if len(binary_cols) >= 2:
        charts.append({
            "type": "multibar",
            "title": "항목별 현황",
            "cols": [
                {"col": c["key"], "label": c["label"].replace("O_", "")}
                for c in binary_cols
            ],
        })

    # category 차트가 부족하면 일반 numeric 컬럼으로 bar 차트 보충 (TS buildDefaultConfig 동기화)
    plain_nums = [c for c in num_cols if not c["key"].startswith("O_")]
    if len(charts) < 4:
        for col in plain_nums[: 4 - len(charts)]:
            charts.append({"col": col["key"], "type": "bar", "title": col["label"]})

    return {
        "version": 1,
        "kpi": kpi[:4],
        "charts": charts[:6],
        "list": {
            "visible_cols": [c["key"] for c in columns[:8]],
            "filter_cols":  [c["key"] for c in cat_cols[:3]],
        },
    }


def _read_cleaned_sheet(cleaned_xlsx: Path, cfg: SurveyConfig) -> tuple[list[str], list[dict]]:
    """xlsx 의 Cleaned 시트 → (headers, rows).

    **이 모듈에서 xlsx 를 아는 유일한 함수다.** 여기서 아래로는 rows(list[dict])만 흐른다.
    """
    # read_only=True is more tolerant of slicer-patched XMLs; fall back if needed
    try:
        wb = openpyxl.load_workbook(cleaned_xlsx, data_only=True, read_only=True)
    except Exception:
        wb = openpyxl.load_workbook(cleaned_xlsx, data_only=True)

    # read_only 워크북은 내부적으로 zip 아카이브 핸들을 계속 들고 있어, 명시적으로
    # close() 하지 않으면 CPython의 순환 참조 때문에 GC가 지연되어 파일 핸들이
    # 예측 불가능하게 오래 열려 있게 된다. 이후 파이프라인 재실행 시 같은 파일을
    # 교체(rename)하려다 WinError 5(액세스 거부)로 실패하는 원인이었다.
    try:
        ws_name = cfg.sheets.cleaned
        if ws_name not in wb.sheetnames:
            raise ValueError(f"'{ws_name}' 시트를 찾을 수 없습니다: {cleaned_xlsx.name}")
        ws = wb[ws_name]

        # ── headers from row 2 (output header row) ───────────────────────────────
        # read_only worksheets use .rows iterator; cell() is not available
        all_rows_iter = list(ws.rows)
    finally:
        wb.close()

    hdr_row_cells = all_rows_iter[1] if len(all_rows_iter) > 1 else []
    headers: list[str] = [
        str(c.value).strip() if c.value is not None else ""
        for c in hdr_row_cells
    ]

    # ── rows from row 3+ ──────────────────────────────────────────────────────
    rows: list[dict] = []
    for row_cells in all_rows_iter[2:]:
        row: dict = {}
        for ci, (h, cell) in enumerate(zip(headers, row_cells), 1):
            if not h:
                continue
            row[h] = cell.value
        if any(v is not None and str(v).strip() != "" for v in row.values()):
            rows.append(row)

    return headers, rows


def _headers_from_rows(rows: list[dict]) -> list[str]:
    """rows 에서 컬럼 순서를 복원한다 (첫 등장 순서 유지).

    편집 경로처럼 headers 를 따로 들고 있지 않은 호출부용. dict 는 삽입 순서를 보존하므로
    data.json 을 읽어 그대로 넘기면 원래 컬럼 순서가 유지된다.
    """
    seen: dict[str, None] = {}
    for row in rows:
        for k in row:
            seen[k] = None
    return list(seen)


def clean_value(v: Any) -> Any:
    """셀 값 → JSON 직렬화 가능한 값. (구 export_to_json 내부의 중첩 함수 `_clean`. overrides 도 쓴다)

    셀 하나당 1회 호출된다(15k행 × 39열이면 60만 회) — 모듈 수준에 두고 import 를 밖으로
    올린 이유다. 동작은 분리 전과 동일하다.
    """
    if v is None:
        return None
    # NaN 값 검사 및 방어
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, (np.integer, np.floating)) and np.isnan(v):
        return None

    if isinstance(v, (int, float, bool, np.integer, np.floating)):
        return float(v) if isinstance(v, (float, np.floating)) else int(v)
    val_str = str(v).strip()
    if val_str == "NaN" or val_str == "nan":
        return None
    if val_str.isdigit():
        return int(val_str)
    try:
        val_float = float(val_str)
        if math.isnan(val_float):
            return None
        return val_float
    except ValueError:
        pass
    return val_str


def column_types_of(data: dict) -> dict[str, str]:
    """기존 data.json → {컬럼키: 타입}. 편집 경로가 `column_types` 로 되먹일 값."""
    return {c["key"]: c["type"] for c in data.get("meta", {}).get("columns", [])}


def build_data_json(
    rows: list[dict],
    cfg: SurveyConfig,
    *,
    headers: list[str] | None = None,
    column_types: dict[str, str] | None = None,
    rows_are_clean: bool = False,
    project_dir: Path | None = None,
    source_file: str = "",
) -> dict:
    """rows → 웹 대시보드용 JSON dict. **xlsx 를 열지 않는다.**

    `aggregates`·`unique_values`·min/max/sum 을 계산하는 유일한 지점이다.
    두 경로가 이 함수를 공유하므로 산출물이 어긋날 수 없다:
      - `/run`·`/export`  : xlsx 읽기 → rows → 여기 (타입을 감지한다)
      - 편집 저장          : data.json 읽기 → rows 패치 → 여기 (타입을 물려받는다)

    headers: 컬럼 순서. None 이면 rows 에서 복원한다(`_headers_from_rows`).
    column_types: **편집 경로 전용.** {컬럼키: 타입} 을 주면 그 컬럼은 타입을 재감지하지 않는다.
        왜 필요한가 — 타입은 xlsx **원시값**으로 감지되는데(텍스트 서식의 코드 컬럼 "3000000"
        → text), rows 에는 `clean_value` 를 거친 값이 실린다(→ 3000000). 그래서 data.json 의
        rows 로 재감지하면 text 가 numeric 으로 **뒤집힌다.** 값 하나 고쳤다고 컬럼의 성격이
        바뀌면 안 된다 — **편집은 값을 바꾸는 것이지 타입을 바꾸는 것이 아니다.**
        타입만 물려받고 `unique_values`·min/max/sum 은 현재 값으로 다시 계산한다
        (그래야 편집으로 생긴 새 값이 드로어 드롭다운에 나온다).
        모르는 컬럼은 평소대로 감지하므로, 컬럼이 새로 생겨도 안전하다.
        상세: docs/plan/pending/dashboard_edit_plan.md §5.1 ⑥
    rows_are_clean: **편집 경로 전용.** rows 가 이미 `clean_value` 를 거친 값이면 True.
        그러면 전 셀 재정제(sangga 기준 60만 회·493ms)를 건너뛴다 — 멱등이라 결과가 같으므로
        순수한 낭비다. 대신 **호출자가 새로 넣은 값은 직접 clean_value 를 거쳐야 한다.**
        xlsx 경로는 openpyxl 원시값을 주므로 False(기본)여야 한다.
    project_dir: dashboard.json 이 있으면 result["dashboard"] 에 포함.
    source_file: meta.source_file 에 들어갈 이름. 편집 경로는 xlsx 가 없으므로 빈 문자열 허용.
    """
    if headers is None:
        headers = _headers_from_rows(rows)

    # ── column metadata ───────────────────────────────────────────────────────
    # __row_id 는 rows 에는 실려 나가지만 meta.columns 에는 넣지 않는다. 프런트는 컬럼 목록을
    # meta.columns 로 그리므로(표·차트·필터·드로어·CSV 내보내기) 여기서 빼는 것만으로
    # 사용자 눈에 띄지 않는다. 타입 감지·aggregates 대상에서도 자연히 제외된다.
    columns: list[dict] = []
    for h in headers:
        if not h or h == ROW_ID_COL:
            continue
        vals = [r[h] for r in rows if r.get(h) is not None]
        # 타입은 물려받은 게 있으면 그걸 쓰고(편집 경로), 없으면 감지한다(xlsx 경로).
        # 나머지 통계는 어느 경우든 **현재 값**으로 다시 계산한다.
        col_type = (column_types or {}).get(h) or _detect_type(vals)
        unique_strs = sorted({str(v).strip() for v in vals if str(v).strip()}) if col_type == "category" else []
        entry: dict = {
            "key":          h,
            "label":        h,
            "type":         col_type,
        }
        if col_type == "category":
            entry["unique_count"] = len(unique_strs)
            entry["unique_values"] = unique_strs[:60]   # cap for JSON size
        elif col_type == "numeric":
            nums = [float(v) for v in vals if isinstance(v, (int, float))]
            if nums:
                entry["min"] = min(nums)
                entry["max"] = max(nums)
                entry["sum"] = sum(nums)
            # 연도처럼 정수값이지만 실질적으로는 이산 카테고리인 numeric 컬럼을
            # 프론트(ChartCard)가 구분할 수 있도록 고유값 개수가 적을 때만 부가.
            distinct = {str(v).strip() for v in vals if str(v).strip()}
            if 0 < len(distinct) <= 40:
                entry["unique_count"] = len(distinct)
        columns.append(entry)

    # ── dashboard config 조기 로드 (aggregates 계산에 filter_cols 필요) ────────
    # 우선순위: project/dashboard.json > config/dashboard_defaults.yaml > 자동 생성
    # 자동 생성 결과는 React DashboardConfig 타입과 호환됩니다.
    dash_cfg: dict | None = None

    if project_dir is not None:
        defaults_path = project_dir.parent.parent / "config" / "dashboard_defaults.yaml"
        if defaults_path.exists():
            import yaml as _yaml
            with open(defaults_path, encoding="utf-8") as f:
                raw_defaults = _yaml.safe_load(f) or {}
            if raw_defaults and any(v is not None for v in raw_defaults.values()):
                dash_cfg = raw_defaults

    if project_dir is not None:
        dash_path = project_dir / "dashboard.json"
        if dash_path.exists():
            with open(dash_path, encoding="utf-8") as f:
                dash_cfg = json.load(f)
            print(f"대시보드 설정 포함: {dash_path}")

    if not dash_cfg:
        dash_cfg = _build_default_dashboard(columns)

    filter_cols: set[str] = set(dash_cfg.get("list", {}).get("filter_cols") or [])

    # ── aggregates (category counts) ──────────────────────────────────────────
    # 필터 드롭다운 옵션 개수 상한. 사용자가 필터로 명시 등록한 컬럼이라도
    # 값이 사실상 자유 텍스트(고유값이 지나치게 많음)면 드롭다운이 무의미하고
    # JSON 용량만 커지므로 상한을 넘으면 건너뛴다.
    _FILTER_UNIQUE_CAP = 150
    aggregates: dict[str, dict[str, int]] = {}
    for col in columns:
        vals = [r[col["key"]] for r in rows if r.get(col["key"]) is not None]
        unique_vals = {str(v).strip() for v in vals if str(v).strip()}

        # Check if the column is configured to be in slicer in config.yaml
        cfg_col = next((c for c in cfg.columns if c.output_col == col["key"]), None)
        in_slicer = getattr(cfg_col, "include_in_slicer", False) if cfg_col else False
        is_filter_col = col["key"] in filter_cols

        if is_filter_col and len(unique_vals) > _FILTER_UNIQUE_CAP:
            continue

        if col["type"] == "category" or in_slicer or is_filter_col:
            counts: dict[str, int] = {}
            for row in rows:
                v = str(row.get(col["key"]) or "").strip()
                if v:
                    # 쉼표 분리 후 개별 항목 카운팅 (다중 응답 필터 지원)
                    parts = [p.strip() for p in v.split(",") if p.strip()]
                    for p in parts:
                        counts[p] = counts.get(p, 0) + 1
            aggregates[col["key"]] = dict(sorted(counts.items(), key=lambda x: -x[1]))

    # ── serialize rows (convert non-JSON types) ───────────────────────────────
    # 이미 정제된 rows 를 다시 정제하는 건 멱등이라 결과가 같다 — 편집 경로에서는 그 한 번이
    # 60만 회 호출(15k행 × 39열)이라 순수한 비용이다. 그래서 건너뛸 수 있게 열어둔다.
    if rows_are_clean:
        clean_rows = [{h: row.get(h) for h in headers if h} for row in rows]
    else:
        clean_rows = [{h: clean_value(row.get(h)) for h in headers if h} for row in rows]

    result = {
        "meta": {
            "project":      cfg.project,
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "total_rows":   len(rows),
            "source_file":  source_file,
            "columns":      columns,
        },
        "rows":       clean_rows,
        "aggregates": aggregates,
    }

    # dash_cfg는 aggregates 계산 전에 이미 로드됨 (filter_cols 참조 목적)
    result["dashboard"] = dash_cfg

    return result


def write_data_json(result: dict, output_path: Path, cfg: SurveyConfig) -> None:
    """build_data_json() 결과를 파일로 쓰고 manifest 를 갱신한다."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        # 프론트엔드가 그대로 fetch해서 파싱하는 정적 데이터 파일이라 사람이 읽을
        # 필요가 없다. indent + 여분 공백을 없애면 파일 크기가 크게 줄어든다
        # (15k행 데이터셋 기준 약 15~20%).
        json.dump(result, f, ensure_ascii=False, separators=(",", ":"))
    print(
        f"내보내기 완료: {output_path} "
        f"({result['meta']['total_rows']}행, {len(result['meta']['columns'])}컬럼)"
    )
    _update_manifest(output_path, cfg)


def export_to_json(
    cleaned_xlsx: Path,
    cfg: SurveyConfig,
    output_path: Path | None = None,
    project_dir: Path | None = None,
) -> dict:
    """Read Cleaned sheet from xlsx → return JSON-serializable dict.

    project_dir: 프로젝트 폴더 경로. dashboard.json이 있으면 result["dashboard"]에 포함.
    Also writes to output_path if provided.

    (2026-07-17) 내부가 `_read_cleaned_sheet` + `build_data_json` + `write_data_json` 으로
    분리됐다. 시그니처·동작·산출물은 이전과 동일하다.
    """
    headers, rows = _read_cleaned_sheet(cleaned_xlsx, cfg)
    result = build_data_json(
        rows,
        cfg,
        headers=headers,
        project_dir=project_dir,
        source_file=cleaned_xlsx.name,
    )

    if output_path:
        write_data_json(result, output_path, cfg)

    return result


def _update_manifest(data_json: Path, cfg: SurveyConfig) -> None:
    """Update frontend/public/data/projects.json manifest."""
    manifest_path = data_json.parent / "projects.json"
    existing: list[dict] = []
    if manifest_path.exists():
        try:
            with open(manifest_path, encoding="utf-8") as f:
                existing = json.load(f)
        except Exception:
            existing = []

    entry = {
        "id":      cfg.project,
        "name":    cfg.project,
        "file":    data_json.name,
        "updated": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    existing = [e for e in existing if e.get("id") != cfg.project]
    existing.append(entry)

    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(existing, f, ensure_ascii=False, indent=2)
    print(f"프로젝트 목록 갱신: {manifest_path}")
