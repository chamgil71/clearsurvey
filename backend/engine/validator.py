from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from engine.config import SummarySection, SurveyConfig
from engine.patterns import load_patterns
from engine.writer import expanded_output_col_names


@dataclass
class ValidationReport:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    infos: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def error(self, message: str) -> None:
        self.errors.append(message)

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def info(self, message: str) -> None:
        self.infos.append(message)


def _resolve_project_dir(config_path: Path | None) -> Path | None:
    if config_path is None:
        return None
    if config_path.suffix.lower() in (".yaml", ".yml"):
        return config_path.parent
    return config_path.parent.parent if config_path.parent.name.lower() == "output" else config_path.parent


def _resolve_source_path(
    cfg: SurveyConfig,
    config_path: Path | None,
    input_path: Path | None,
) -> Path | None:
    if input_path is not None:
        return input_path
    proj_dir = _resolve_project_dir(config_path)
    if cfg.source.file and proj_dir:
        src = Path(cfg.source.file)
        resolved = src if src.is_absolute() else proj_dir / src
        if not resolved.exists():
            fname = src.name
            fallback_1 = proj_dir.parent.parent / "raw" / fname
            if fallback_1.exists():
                return fallback_1
            fallback_2 = Path("c:/ai/clearsurvey/storage/raw") / fname
            if fallback_2.exists():
                return fallback_2
        return resolved
    return None


def _source_column_count(
    cfg: SurveyConfig,
    config_path: Path | None,
    input_path: Path | None,
    report: ValidationReport,
) -> int | None:
    if cfg.merge:
        report.info("merge 설정 사용: 원본 컬럼 범위 검사는 생략합니다.")
        return None

    src_path = _resolve_source_path(cfg, config_path, input_path)
    if src_path is None:
        report.warn("원본 파일 경로를 알 수 없어 source_col 범위 검사를 생략합니다. --input 옵션을 사용하면 검사할 수 있습니다.")
        return None
    if not src_path.exists():
        report.error(f"원본 파일을 찾을 수 없습니다: {src_path}")
        return None

    try:
        wb = openpyxl.load_workbook(src_path, read_only=True, data_only=True)
    except Exception as exc:
        report.error(f"원본 파일을 열 수 없습니다: {src_path} ({exc})")
        return None

    try:
        sheet_name = cfg.source.sheet
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        elif sheet_name:
            report.error(f"원본 시트를 찾을 수 없습니다: {sheet_name}")
            return None
        else:
            ws = wb.active
        report.info(f"원본 컬럼 수: {ws.max_column}")
        return int(ws.max_column)
    finally:
        wb.close()


def _prepare_registry(cfg: SurveyConfig, config_path: Path | None):
    """Build the same transform registry shape used by SurveyPipeline."""
    from engine.config import AddressParsingConfig
    from engine.pipeline import _build_registry

    proj_dir = _resolve_project_dir(config_path)
    patterns = load_patterns(cfg.patterns_file, proj_dir)
    if cfg.address_parsing is None and patterns.address_parsing:
        cfg.address_parsing = AddressParsingConfig.model_validate(patterns.address_parsing)
    return _build_registry(cfg)


def _check_positive_col(report: ValidationReport, label: str, value: int | None) -> None:
    if value is not None and value < 1:
        report.error(f"{label}은 1 이상의 정수여야 합니다: {value}")


def _check_col_bound(report: ValidationReport, label: str, value: int | None, max_col: int | None) -> None:
    _check_positive_col(report, label, value)
    if value is not None and max_col is not None and value > max_col:
        report.error(f"{label}이 원본 컬럼 수({max_col})를 초과합니다: {value}")


def _section_refs(sec: SummarySection) -> list[tuple[str, str | None]]:
    refs: list[tuple[str, str | None]] = []
    if sec.col_ref:
        refs.append((f"summary[{sec.id or sec.title}].col_ref", sec.col_ref))
    for i, item in enumerate(sec.items, 1):
        refs.append((f"summary[{sec.id or sec.title}].items[{i}].col_ref", item.col_ref))
    for i, item in enumerate(sec.columns, 1):
        refs.append((f"summary[{sec.id or sec.title}].columns[{i}].col_ref", item.col_ref))
    return refs


def validate_config(
    cfg: SurveyConfig,
    *,
    config_path: Path | None = None,
    input_path: Path | None = None,
) -> ValidationReport:
    report = ValidationReport()

    registry = _prepare_registry(cfg, config_path)
    max_col = _source_column_count(cfg, config_path, input_path, report)

    active_columns = [col for col in cfg.columns if col.transform != "exclude"]
    if not active_columns:
        report.error("출력 컬럼이 비어 있습니다. transform=exclude가 아닌 컬럼이 최소 1개 필요합니다.")
    output_cols = [col.output_col for col in active_columns]
    output_set = set(output_cols)
    if len(output_cols) != len(output_set):
        seen: set[str] = set()
        dups: set[str] = set()
        for name in output_cols:
            if name in seen:
                dups.add(name)
            seen.add(name)
        report.error(f"output_col 중복: {', '.join(dups)}")

    # addr_split/norm_date_parts/split_binary 등 파생열 확장 후의 실제 출력
    # 컬럼명까지 포함해 충돌을 검사한다. output_col 기본 이름만 검사하면
    # 예: "생년월일"(norm_date_parts, 파생열 "생년월일_년" 자동 생성)과
    # 별도 컬럼 "생년월일_년"이 조용히 같은 이름으로 충돌해도 잡히지 않는다.
    expanded_names = expanded_output_col_names(active_columns)
    if len(expanded_names) != len(set(expanded_names)):
        seen_expanded: set[str] = set()
        expanded_dups: set[str] = set()
        for name in expanded_names:
            if name in seen_expanded:
                expanded_dups.add(name)
            seen_expanded.add(name)
        report.error(
            "파생열 확장 후 출력 컬럼명이 충돌합니다: "
            f"{', '.join(sorted(expanded_dups))} "
            "(addr_split/norm_date_parts/split_binary가 자동 생성하는 파생열 이름이 "
            "다른 컬럼의 output_col과 같습니다. output_col을 변경하세요.)"
        )

    fill_labels = {rule.output_label for rule in cfg.preprocess.fill_down}
    for i, rule in enumerate(cfg.preprocess.fill_down, 1):
        _check_col_bound(report, f"preprocess.fill_down[{i}].source_col", rule.source_col, max_col)
        _check_col_bound(report, f"preprocess.fill_down[{i}].trigger_col", rule.trigger_col, max_col)

    for group_name, rules in (
        ("include", cfg.preprocess.row_filter.include),
        ("exclude", cfg.preprocess.row_filter.exclude),
    ):
        for i, rule in enumerate(rules, 1):
            _check_col_bound(report, f"preprocess.row_filter.{group_name}[{i}].col", rule.col, max_col)

    _ADDRESS_TRANSFORMS = {"address_sido", "address_sigungu", "addr_split"}

    for i, col in enumerate(cfg.columns, 1):
        prefix = f"columns[{i}]({col.output_col})"
        if col.transform == "exclude":
            report.info(f"{prefix}은 transform=exclude로 출력에서 제외됩니다.")
            continue
        if col.transform and col.transform not in registry:
            if col.transform in _ADDRESS_TRANSFORMS:
                report.error(
                    f"{prefix}.transform '{col.transform}' 사용 시 "
                    "config/patterns.yaml의 address_parsing 섹션 설정이 필요합니다.\n"
                    "         → patterns.yaml에 address_parsing.sido_patterns 를 추가하거나 "
                    "config.yaml에 address_parsing 섹션을 추가하세요."
                )
            else:
                report.error(f"{prefix}.transform이 등록되어 있지 않습니다: {col.transform}")
        if col.source_label and col.source_label not in fill_labels:
            report.error(f"{prefix}.source_label이 preprocess.fill_down output_label에 없습니다: {col.source_label}")
        _check_col_bound(report, f"{prefix}.source_col", col.source_col, max_col)
        _check_col_bound(report, f"{prefix}.backup_col", col.backup_col, max_col)
        _check_col_bound(report, f"{prefix}.year_col", col.year_col, max_col)
        if col.source_cols:
            for j, source_col in enumerate(col.source_cols, 1):
                _check_col_bound(report, f"{prefix}.source_cols[{j}]", source_col, max_col)

    for i, slicer in enumerate(cfg.slicers, 1):
        if slicer.col not in output_set:
            report.error(f"slicers[{i}].col이 출력 컬럼에 없습니다: {slicer.col}")
        else:
            col_def = next((col for col in cfg.columns if col.output_col == slicer.col), None)
            if col_def and not col_def.include_in_slicer:
                report.warn(f"slicers[{i}].col은 include_in_slicer=true가 아닙니다: {slicer.col}")

    layout = cfg.summary.layout
    if layout:
        if layout.cols < 1:
            report.error("summary.layout.cols는 1 이상이어야 합니다.")
        if layout.start_row < 1:
            report.error("summary.layout.start_row는 1 이상이어야 합니다.")
        if layout.col_span < 1:
            report.error("summary.layout.col_span은 1 이상이어야 합니다.")

    for i, sec in enumerate(cfg.summary.sections, 1):
        sec_name = sec.id or sec.title or str(i)
        if not layout and (sec.start_row is None or sec.start_col is None):
            report.error(f"summary.sections[{i}]({sec_name})는 layout이 없으면 start_row/start_col이 필요합니다.")
        if layout and (sec.layout_col < 1 or sec.layout_col > layout.cols):
            report.error(f"summary.sections[{i}]({sec_name}).layout_col이 layout.cols 범위를 벗어납니다: {sec.layout_col}")
        if sec.type == "totals" and not sec.items:
            report.warn(f"summary.sections[{i}]({sec_name}) totals 섹션에 items가 없습니다.")
        if sec.type in ("binary_sum", "gpu_demand") and not sec.columns:
            report.warn(f"summary.sections[{i}]({sec_name}) {sec.type} 섹션에 columns가 없습니다.")
        if sec.type == "countif_contains" and not sec.keywords:
            report.warn(f"summary.sections[{i}]({sec_name}) countif_contains 섹션에 keywords가 없습니다.")
        for ref_label, ref in _section_refs(sec):
            if ref and ref not in output_set:
                report.error(f"{ref_label}이 출력 컬럼에 없습니다: {ref}")

    for path_label, path_value in (
        ("style_file", cfg.style_file),
        ("patterns_file", cfg.patterns_file),
    ):
        if not path_value or config_path is None:
            continue
        proj_dir = _resolve_project_dir(config_path)
        path = Path(path_value)
        resolved = path if path.is_absolute() else (proj_dir / path if proj_dir else path)
        if not resolved.exists():
            report.warn(f"{path_label} 파일을 찾을 수 없습니다: {resolved}")

    if report.ok:
        report.info("설정 참조 검증을 통과했습니다.")
    return report
