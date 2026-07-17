from __future__ import annotations

from typing import TYPE_CHECKING, Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from engine.config import ROW_ID_COL, ColumnDef, SurveyConfig, make_row_id
from engine.styler import build_cleaned_styles, load_style
from transforms.registry import TransformRegistry

if TYPE_CHECKING:                     # 런타임 순환 import 회피
    from engine.overrides import OverrideApplier

MAX_DATA_ROW = 10000

# transforms that should use SUBTOTAL(9,…) = SUM instead of COUNTA
_NUMERIC_ROLES = {"n_jang", "jang", "o_binary", "to_binary", "split_binary", "norm_num", "normalize_number", "to_numeric"}

# transforms that return a dict → multiple derived output columns
# key: transform name  /  value: ordered list of column suffixes
# "" = main column (output_col as-is), others appended to output_col
_DERIVED_SUFFIXES: dict[str, list[str]] = {
    "norm_date_parts": ["", "_년", "_월", "_일"],
    "addr_split":      ["", "_시도", "_시군구", "_상세"],
}


def _suffixes_for(col_def: ColumnDef) -> list[str]:
    """다중 출력 transform의 파생 접미사 목록을 반환한다. "" = 본 컬럼."""
    if col_def.transform == "split_binary" and col_def.flag_keyword:
        kws = [k.strip() for k in col_def.flag_keyword.split(",") if k.strip()]
        return [""] + [f"_{kw}" for kw in kws]
    return _DERIVED_SUFFIXES.get(col_def.transform or "", [""])


def expanded_output_col_names(col_defs: list[ColumnDef]) -> list[str]:
    """파생열 확장 후 실제로 쓰여지는 출력 컬럼명 전체 목록을 반환한다.

    addr_split/norm_date_parts/split_binary처럼 컬럼 하나가 여러 출력 컬럼으로
    확장되는 경우까지 포함하므로, output_col 기본 이름만으로는 잡히지 않는
    파생열 간 이름 충돌(예: 생년월일_년 이라는 별도 컬럼과 생년월일의 파생열
    생년월일_년)을 검증할 때 사용한다.
    """
    return [col_def.output_col + sfx for col_def in col_defs for sfx in _suffixes_for(col_def)]


def _resolve_val(row: pd.Series, col_def: ColumnDef) -> Any:
    if col_def.source_cols is not None:
        return [
            (row.iloc[sc - 1] if sc - 1 < len(row) else None)
            for sc in col_def.source_cols
        ]
    if col_def.source_col is not None:
        pos = col_def.source_col - 1
        return row.iloc[pos] if pos < len(row) else None
    if col_def.source_label is not None:
        return row.get(col_def.source_label)
    return None


def _apply_transform(
    col_def: ColumnDef,
    row: pd.Series,
    registry: TransformRegistry,
    extra_kwargs: dict,
) -> Any:
    raw = _resolve_val(row, col_def)
    if col_def.transform is None or col_def.transform not in registry:
        return raw if raw is not None and str(raw).strip() != "" else None

    kw = dict(extra_kwargs)
    if col_def.flag_keyword is not None:
        kw["flag_keyword"] = col_def.flag_keyword
    if col_def.backup_col is not None:
        bkup_pos = col_def.backup_col - 1
        kw["backup_val"] = row.iloc[bkup_pos] if bkup_pos < len(row) else None
    if col_def.year_col is not None:
        yr_pos = col_def.year_col - 1
        kw["year_val"] = row.iloc[yr_pos] if yr_pos < len(row) else None

    return registry.apply(col_def.transform, raw, **kw)


class CleanedSheetWriter:
    def __init__(self, cfg: SurveyConfig, registry: TransformRegistry):
        self._cfg      = cfg
        self._registry = registry

    def _expand_cols(
        self, col_defs: list[ColumnDef]
    ) -> list[tuple[ColumnDef, str]]:
        """Expand ColumnDefs with derived suffixes for multi-output transforms.

        Returns list of (ColumnDef, suffix) pairs.
        "" suffix = main column (output_col unchanged).
        Non-"" suffix is appended to output_col name.
        """
        expanded: list[tuple[ColumnDef, str]] = []
        for cd in col_defs:
            for sfx in _suffixes_for(cd):
                expanded.append((cd, sfx))
        return expanded

    def write(
        self,
        wb: Workbook,
        df: pd.DataFrame,
        applier: "OverrideApplier | None" = None,
    ) -> tuple[int, dict[str, int], dict[str, list]]:
        """Write Cleaned sheet into wb from preprocessed DataFrame.

        Returns (row_count, col_index_map, cleaned_col_vals).
        col_index_map: {output_col_name: 1-based column index} — includes derived cols.

        applier: 손 편집 오버레이(engine/overrides.py). transform 결과를 시트에 쓰기
        **직전**에 통과시킨다 — 정제 규칙과 손 편집이 부딪히면 손 편집이 이긴다.
        None 이면 아무것도 덮어쓰지 않는다(기존 동작).
        """
        style = load_style(None, self._cfg.style_file)
        s     = build_cleaned_styles(style)

        sheet_name = self._cfg.sheets.cleaned
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)

        base_defs = [cd for cd in self._cfg.columns if cd.transform != "exclude"]
        expanded  = self._expand_cols(base_defs)   # [(ColumnDef, suffix), ...]
        n_cols    = len(expanded)

        # ── __row_id: 데이터 컬럼 뒤에 붙는 숨김 열 ──────────────────────────
        # 사용자 컬럼과 이름이 겹치면 내보내기 시 dict 키가 충돌해 한쪽이 조용히 사라진다.
        # 조용한 데이터 손상보다 즉시 실패가 낫다.
        if any(cd.output_col + sfx == ROW_ID_COL for (cd, sfx) in expanded):
            raise ValueError(
                f"출력 컬럼명 '{ROW_ID_COL}' 은 행 식별자 예약어라 사용할 수 없습니다. "
                f"config.yaml 의 output_col 을 다른 이름으로 바꾸세요."
            )
        row_id_ci   = n_cols + 1
        last_letter = get_column_letter(row_id_ci)

        extra_kw = dict(self._cfg.transform_kwargs)
        if self._cfg.jang_extraction:
            extra_kw["jang_cfg"] = self._cfg.jang_extraction.model_dump()

        # ── Row 1: SUBTOTAL aggregation row ─────────────────────────────────
        for ci, (cd, sfx) in enumerate(expanded, 1):
            letter = get_column_letter(ci)
            rng    = f"{letter}3:{letter}{MAX_DATA_ROW}"
            fn     = 9 if cd.transform in _NUMERIC_ROLES else 3
            c = ws.cell(1, ci, f"=SUBTOTAL({fn},{rng})")
            c.fill, c.font, c.alignment, c.border = (
                s["sub_fill"], s["sub_font"], s["sub_align"], s["sub_bdr"]
            )
        # __row_id 열은 집계가 무의미하므로 수식 없이 서식만 맞춘다(행 1의 띠가 끊기지 않도록).
        c = ws.cell(1, row_id_ci)
        c.fill, c.font, c.alignment, c.border = (
            s["sub_fill"], s["sub_font"], s["sub_align"], s["sub_bdr"]
        )
        ws.row_dimensions[1].height = s["sub_height"]

        # ── Row 2: header labels ─────────────────────────────────────────────
        for ci, (cd, sfx) in enumerate(expanded, 1):
            header = cd.output_col + sfx
            c = ws.cell(2, ci, header)
            c.fill, c.font, c.alignment = s["hdr_fill"], s["hdr_font"], s["hdr_align"]
        c = ws.cell(2, row_id_ci, ROW_ID_COL)
        c.fill, c.font, c.alignment = s["hdr_fill"], s["hdr_font"], s["hdr_align"]
        ws.row_dimensions[2].height = s["hdr_height"]

        # ── Rows 3+: data ────────────────────────────────────────────────────
        row_out = 3
        n_total = len(df)
        _PROG   = 1000
        cleaned_col_vals: dict[str, list] = {
            cd.output_col + sfx: [] for (cd, sfx) in expanded
        }
        for ri, (row_id, row) in enumerate(df.iterrows()):
            if n_total >= _PROG and ri > 0 and ri % _PROG == 0:
                pct = int(ri / n_total * 100)
                print(f"  변환중: {ri:,}/{n_total:,}행 ({pct}%) ...", end="\r", flush=True)

            rid = make_row_id(row_id)
            ci = 1
            for cd in base_defs:
                val = _apply_transform(cd, row, self._registry, extra_kw)
                if cd.transform == "split_binary" and cd.flag_keyword:
                    kws = [k.strip() for k in cd.flag_keyword.split(",") if k.strip()]
                    suffixes = [""] + [f"_{kw}" for kw in kws]
                else:
                    suffixes = _DERIVED_SUFFIXES.get(cd.transform or "", [""])
                for sfx in suffixes:
                    if isinstance(val, dict):
                        cell_val = val.get(sfx)
                    elif sfx == "":
                        cell_val = val
                    else:
                        cell_val = None
                    col_key = cd.output_col + sfx
                    # ★ 손 편집 덮어쓰기 지점 — transform 이후, 시트 기록 직전.
                    # cleaned_col_vals 에 담기 전에 통과시켜야 Summary 시트의 엑셀 차트도
                    # 편집을 반영한다(요약 집계가 이 값을 쓴다).
                    if applier is not None:
                        cell_val = applier.take(rid, col_key, cell_val)
                    cleaned_col_vals[col_key].append(cell_val)
                    cell = ws.cell(row_out, ci, cell_val)
                    cell.font      = s["dat_font"]
                    cell.alignment = s["dat_align"]
                    if cd.number_format and sfx == "":
                        cell.number_format = cd.number_format
                    if cd.align and sfx == "":
                        from openpyxl.styles import Alignment as _Align
                        cell.alignment = _Align(horizontal=cd.align, vertical="center")
                    ci += 1
            # 행 식별자 — df.index 가 나른 원본 엑셀 행번호 (병합 경로는 0-based 위치)
            rid_cell = ws.cell(row_out, row_id_ci, rid)
            rid_cell.font      = s["dat_font"]
            rid_cell.alignment = s["dat_align"]
            row_out += 1

        if n_total >= _PROG:
            print(f"  변환중: {n_total:,}/{n_total:,}행 (100%) 완료    ")

        last_row = max(row_out - 1, 2)

        # ── Excel Table ──────────────────────────────────────────────────────
        table_ref = f"A2:{last_letter}{last_row}"
        tab = Table(displayName="Survey", ref=table_ref)
        tab.tableStyleInfo = s["tbl_style"]
        ws.add_table(tab)

        # ── column widths ────────────────────────────────────────────────────
        for ci, (cd, sfx) in enumerate(expanded, 1):
            ws.column_dimensions[get_column_letter(ci)].width = cd.width or 14

        # __row_id 는 사용자에게 보일 이유가 없다 — 숨긴다.
        # 다만 Excel Table(A2:{last_letter}) **안에** 포함시켜, 사용자가 엑셀에서 정렬·필터해도
        # 식별자가 제 행에 붙어 다니게 한다. 수정된 xlsx 를 되돌려 읽을 때 행 매칭의 근거다.
        rid_letter = get_column_letter(row_id_ci)
        ws.column_dimensions[rid_letter].width  = 12
        ws.column_dimensions[rid_letter].hidden = True

        ws.freeze_panes = s["freeze_pane"]

        # col_index_map 에는 __row_id 를 넣지 않는다 — 이 맵은 슬라이서 배치(slicer.py 의
        # n_data_cols)와 요약 시트가 "실제 데이터 컬럼"을 세는 데 쓰인다. 숨은 식별자가
        # 거기 끼면 슬라이서가 한 칸씩 밀린다.
        col_index_map = {
            cd.output_col + sfx: ci
            for ci, (cd, sfx) in enumerate(expanded, 1)
        }
        return row_out - 3, col_index_map, cleaned_col_vals
