from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine.config import ROW_ID_COL, SurveyConfig, make_row_id
from engine.csv_analyzer import read_csv_rows
from engine.json_analyzer import json_records_to_columns, load_json_records
from engine.overrides import OverrideApplier, load_overrides
from engine.preprocessor import Preprocessor
from engine.writer import CleanedSheetWriter
from engine.summarizer import SummarySheetWriter
from engine.slicer import inject_slicers
from engine.patterns import load_patterns, inject_into_kwargs
from transforms.registry import TransformRegistry


def _sheet_to_dataframe(wb: openpyxl.Workbook, cfg: SurveyConfig) -> pd.DataFrame:
    """Read source sheet into a positionally-indexed DataFrame.

    DataFrame 의 **index 는 원본 엑셀의 실제 행번호**다(`__row_id` 의 출처, engine/config.py 참고).
    빈 행을 건너뛰어도 번호는 원본 기준으로 유지되므로, 이후 원본에 행이 추가·삭제되어도
    살아남은 행의 id 는 변하지 않는다.

    행번호를 **컬럼이 아니라 index 로** 보존하는 이유: `_resolve_val` 이 `row.iloc[source_col-1]`
    로 위치 기반 접근을 하고 `len(row)` 로 경계를 검사한다. 컬럼을 하나 늘리면 그 경계가 밀려
    잘못 설정된 source_col 이 행번호를 데이터로 읽어갈 수 있다. index 는 iloc 에 관여하지 않는다.
    """
    sheet_name  = cfg.source.sheet
    ws          = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    if ws is None:
        raise ValueError("Source sheet not found")

    data_start = cfg.source.data_start_row or (cfg.source.header_row + 1)
    rows: list[list[Any]] = []
    row_ids: list[int] = []
    # iter_rows(values_only=True)로 행 단위 스트리밍 — ws.cell(r, c) 개별 호출 대비
    # 좌표 조회/Cell 객체 생성 오버헤드가 없어 대용량 시트에서 수백 배 빠르다.
    # enumerate(start=data_start): iter_rows 가 내주는 첫 행이 곧 엑셀 data_start 행이다.
    for excel_row, row_vals in enumerate(
        ws.iter_rows(min_row=data_start, values_only=True), start=data_start
    ):
        # skip entirely blank rows — 건너뛰어도 excel_row 는 원본 번호를 유지한다
        if any(v is not None and str(v).strip() != "" for v in row_vals):
            rows.append(list(row_vals))
            row_ids.append(excel_row)

    if not rows:
        return pd.DataFrame()

    n_cols = max(len(r) for r in rows)
    padded = [r + [None] * (n_cols - len(r)) for r in rows]
    # columns: 0, 1, 2, ... (0-based integers) / index: 원본 엑셀 행번호
    return pd.DataFrame(padded, index=pd.Index(row_ids, name=ROW_ID_COL))


# 앞에 0이 붙지 않은 정수/소수만 숫자로 본다("01012345678" 같은 전화번호 문자열이 앞자리 0을
# 잃고 조용히 정수로 바뀌는 사고를 막는다 — "0" 단독이나 "0.5" 같은 소수는 허용).
_CSV_INT_RE = re.compile(r"^-?(0|[1-9]\d*)$")
_CSV_FLOAT_RE = re.compile(r"^-?(0|[1-9]\d*)\.\d+$")


def _coerce_csv_value(raw: str | None) -> Any:
    """CSV는 모든 셀이 문자열이다 — openpyxl이 주는 타입 있는 값(int/float/None)과 맞춘다.

    `_detect_type`(exporter.py)이 `isinstance(v, (int, float))`로만 숫자를 세므로, 이 보정이
    없으면 csv로 올린 숫자 컬럼이 전부 "text"로 오분류된다.
    """
    if raw is None:
        return None
    s = raw.strip()
    if s == "":
        return None
    if _CSV_INT_RE.match(s):
        return int(s)
    if _CSV_FLOAT_RE.match(s):
        return float(s)
    return raw


def _csv_to_dataframe(path: Path, cfg: SurveyConfig) -> pd.DataFrame:
    """CSV를 `_sheet_to_dataframe`과 동일한 계약(0-based 컬럼, 원본 행번호 index)으로 읽는다.

    이 계약만 지키면 `Preprocessor` 이후(transform·override·writer)는 xlsx 경로와 코드 한 줄도
    다르지 않다 — `docs/plan/pending/new_beginnings_comparison_plan.md` §1-B 참고.
    """
    all_rows = read_csv_rows(path)

    data_start = cfg.source.data_start_row or (cfg.source.header_row + 1)
    rows: list[list[Any]] = []
    row_ids: list[int] = []
    for csv_row_num, row_vals in enumerate(all_rows, start=1):
        if csv_row_num < data_start:
            continue
        coerced = [_coerce_csv_value(v) for v in row_vals]
        if any(v is not None and str(v).strip() != "" for v in coerced):
            rows.append(coerced)
            row_ids.append(csv_row_num)

    if not rows:
        return pd.DataFrame()

    n_cols = max(len(r) for r in rows)
    padded = [r + [None] * (n_cols - len(r)) for r in rows]
    return pd.DataFrame(padded, index=pd.Index(row_ids, name=ROW_ID_COL))


def _json_to_dataframe(path: Path, cfg: SurveyConfig) -> pd.DataFrame:
    """배열-of-객체 JSON을 `_sheet_to_dataframe`과 동일 계약으로 읽는다.

    JSON은 이미 값에 타입이 있다(숫자는 정수/실수로, 문자열은 문자열로) — csv처럼 값 보정이
    필요 없다. "행 번호"는 헤더가 가상의 1행을 차지한다고 보고 배열 인덱스+2로 매긴다
    (레코드 0번째 = 2행, 그래야 data_start_row=2 관례와 맞는다).
    """
    records = load_json_records(path)
    _headers, matrix = json_records_to_columns(records)

    data_start = cfg.source.data_start_row or (cfg.source.header_row + 1)
    rows: list[list[Any]] = []
    row_ids: list[int] = []
    for offset, row_vals in enumerate(matrix):
        json_row_num = offset + 2  # 헤더가 가상의 1행
        if json_row_num < data_start:
            continue
        if any(v is not None and str(v).strip() != "" for v in row_vals):
            rows.append(row_vals)
            row_ids.append(json_row_num)

    if not rows:
        return pd.DataFrame()

    n_cols = max(len(r) for r in rows)
    padded = [r + [None] * (n_cols - len(r)) for r in rows]
    return pd.DataFrame(padded, index=pd.Index(row_ids, name=ROW_ID_COL))


def _write_raw_sheet(out_wb: Workbook, src_wb: openpyxl.Workbook | None,
                     cfg: SurveyConfig, raw_df: pd.DataFrame,
                     source_headers: list[str] | None = None) -> None:
    """Write raw (pre-transform) source data to '원본' sheet."""
    sheet_name = "원본"
    if sheet_name in out_wb.sheetnames:
        del out_wb[sheet_name]
    ws = out_wb.create_sheet(sheet_name)

    n_cols = raw_df.shape[1]

    # original header labels from source workbook (xlsx) or csv/json header row
    if src_wb is not None:
        src_ws = (
            src_wb[cfg.source.sheet]
            if cfg.source.sheet and cfg.source.sheet in src_wb.sheetnames
            else src_wb.active
        )
        hrow = cfg.source.header_row
        headers = [src_ws.cell(hrow, c).value for c in range(1, n_cols + 1)]
    elif source_headers is not None:
        headers = [
            source_headers[i] if i < len(source_headers) else f"열{i + 1}"
            for i in range(n_cols)
        ]
    else:
        headers = [f"열{i}" for i in range(1, n_cols + 1)]

    hdr_fill = PatternFill("solid", start_color="1F4E79")
    hdr_font = Font(bold=True, name="Arial", size=9, color="FFFFFF")
    hdr_align = Alignment(horizontal="center", vertical="center")
    dat_font  = Font(name="Arial", size=9)

    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h)
        c.font, c.fill, c.alignment = hdr_font, hdr_fill, hdr_align
    ws.row_dimensions[1].height = 20

    for ri, (_, row) in enumerate(raw_df.iterrows(), 2):
        for ci, val in enumerate(row, 1):
            ws.cell(ri, ci, val).font = dat_font

    for ci in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 14
    ws.freeze_panes = "A2"


def enrich_config_with_patterns(cfg: SurveyConfig, proj_dir: Path | None) -> None:
    """patterns_file 기반 공용 설정(address_parsing, transform_kwargs)을 cfg에 주입한다.

    run()과 미리보기(preview) 양쪽 경로가 이 함수를 공유해야 한다. 그렇지 않으면
    미리보기가 address_parsing 없이 addr_split을 실행해 항상 빈 값을 반환하는
    문제가 재발한다 — 실제로 발생했던 회귀 버그다.
    """
    patterns = load_patterns(cfg.patterns_file, proj_dir)

    # address_parsing: 프로젝트 config.yaml 우선, 없으면 patterns_file 사용
    if cfg.address_parsing is None and patterns.address_parsing:
        from engine.config import AddressParsingConfig
        cfg.address_parsing = AddressParsingConfig.model_validate(
            patterns.address_parsing
        )
        if cfg.address_parsing.sido_patterns:
            print(f"주소 패턴 로드: {cfg.patterns_file} ({len(cfg.address_parsing.sido_patterns)}개 시도)")

    # company/phone/brn 패턴을 transform_kwargs에 주입
    inject_into_kwargs(cfg.transform_kwargs, patterns)


def _build_registry(cfg: SurveyConfig) -> TransformRegistry:
    """파이프라인 실행마다 독립된 TransformRegistry 인스턴스를 생성합니다.

    모듈 수준 싱글톤을 공유하지 않으므로 FastAPI 멀티스레드 환경에서도
    요청 간 transform 등록이 서로 간섭하지 않습니다.
    """
    reg = TransformRegistry()  # ← 글로벌 singleton 대신 요청별 인스턴스 생성

    # ── 내장 pass-through ─────────────────────────────────────────────────────
    reg.register("copy", lambda val, **kw: val)

    # ── 공통 주소 모듈 폴백 (기본값) ──────────────────────────────────────────
    try:
        from transforms.common.address import _TRANSFORMS as _addr_transforms
        reg.register_dict(_addr_transforms)
    except Exception:
        pass

    # ── 주소 변환 (project address_parsing 설정 기반 클로저) ──────────────────
    # 클로저가 해당 프로젝트의 address_parsing 설정을 캡처하므로
    # 싱글톤에 등록하면 동시 요청 간 설정이 덮어써집니다.
    if cfg.address_parsing:
        from transforms.common.address import AddressParser, build_addr_parts_dict
        ap     = AddressParser(cfg.address_parsing.model_dump())
        _cache: dict[str, tuple] = {}

        def _parse_cached(val: object) -> tuple:
            k = str(val) if val is not None else ""
            if k not in _cache:
                _cache[k] = ap.parse(val)
            return _cache[k]

        def _sido(val: object, **kw: object) -> str | None:
            return _parse_cached(val)[0] or None

        def _sigungu(val: object, **kw: object) -> str | None:
            return _parse_cached(val)[1] or None

        def _addr_split(val: object, **kw: object) -> dict | None:
            if not val:
                return None
            sido, sigungu, detail = _parse_cached(val)
            return build_addr_parts_dict(val, sido, sigungu, detail)

        reg.register("address_sido",    _sido)
        reg.register("address_sigungu", _sigungu)
        reg.register("addr_split",      _addr_split)

    # ── domain 모듈 자동 로드 (cleansing.py, gpu_survey.py 등) ───────────────
    try:
        reg.auto_load_domain("transforms.domain")
    except Exception as exc:
        print(f"[경고] domain 모듈 자동 로드 실패: {exc}")

    return reg


class SurveyPipeline:
    def __init__(self, cfg: SurveyConfig, config_path: str | Path | None = None):
        self._cfg         = cfg
        self._config_path = Path(config_path) if config_path else None

    def run(self, input_path: str | Path | None = None, dry_run: bool = False) -> Path | None:
        cfg = self._cfg
        proj_dir = self._config_path.parent if self._config_path else None

        # ── resolve style_file path relative to config ───────────────────────
        if cfg.style_file and proj_dir:
            sf = Path(cfg.style_file)
            if not sf.is_absolute():
                cfg.style_file = str(proj_dir / sf)

        # ── load shared patterns (patterns_file) ──────────────────────────────
        enrich_config_with_patterns(cfg, proj_dir)

        # ── input: merger or direct file ──────────────────────────────────────
        src_wb: openpyxl.Workbook | None = None
        source_headers: list[str] | None = None
        if cfg.merge:
            from engine.merger import DataMerger
            df = DataMerger(cfg.merge).run()
            wb = Workbook()
        else:
            src_path = (
                Path(input_path) if input_path
                else (proj_dir / cfg.source.file if cfg.source.file and proj_dir else None)
            )
            if not src_path or not src_path.exists():
                if cfg.source.file:
                    fname = Path(cfg.source.file).name
                    if proj_dir:
                        fallback_1 = proj_dir.parent.parent / "raw" / fname
                        if fallback_1.exists():
                            src_path = fallback_1
                    if not src_path or not src_path.exists():
                        fallback_2 = Path("c:/ai/clearsurvey/storage/raw") / fname
                        if fallback_2.exists():
                            src_path = fallback_2

            if not src_path or not src_path.exists():
                raise FileNotFoundError(f"입력 파일을 찾을 수 없습니다: {src_path}")
            print(f"읽는 중: {src_path.name}")
            ext = src_path.suffix.lower()
            if ext == ".csv":
                # src_wb는 None으로 남는다 — _write_raw_sheet가 source_headers로 대신 헤더를 얻는다.
                df = _csv_to_dataframe(src_path, cfg)
                all_csv_rows = read_csv_rows(src_path)
                source_headers = all_csv_rows[0] if all_csv_rows else []
            elif ext == ".json":
                df = _json_to_dataframe(src_path, cfg)
                source_headers, _ = json_records_to_columns(load_json_records(src_path))
            else:
                wb = openpyxl.load_workbook(src_path, data_only=True)
                src_wb = wb
                df = _sheet_to_dataframe(wb, cfg)

        print(f"원본 행수: {len(df)}")

        raw_df = df  # keep reference before preprocessing

        # ── preprocess ───────────────────────────────────────────────────────
        if cfg.preprocess.fill_down or cfg.preprocess.row_filter.include or cfg.preprocess.row_filter.exclude:
            df = Preprocessor(cfg.preprocess).run(df)
            print(f"전처리 후 행수: {len(df)}")

        # ── transform registry ────────────────────────────────────────────────
        registry = _build_registry(cfg)

        # ── write sheets ──────────────────────────────────────────────────────
        out_wb = Workbook()
        out_wb.remove(out_wb.active)  # remove default sheet

        # ── 손 편집 오버레이 ──────────────────────────────────────────────────
        # cleaned.xlsx 는 매번 새로 만드는 파생물이라, 대시보드에서 고친 값은 여기서
        # 다시 얹어주지 않으면 이 실행에서 사라진다. overrides.json 이 편집의 진실이다.
        # (docs/plan/pending/dashboard_edit_plan.md §3)
        applier = OverrideApplier(load_overrides(proj_dir)) if proj_dir else None

        writer  = CleanedSheetWriter(cfg, registry)
        n_rows, col_index_map, cleaned_col_vals = writer.write(out_wb, df, applier)
        print(f"Cleaned 시트 기록: {n_rows}행")

        self.override_conflicts: list = []
        if applier is not None and applier.has_edits:
            # 붙을 자리가 사라진 편집은 버리지 않고 보고한다 — 편집이 조용히 없어지는 것이
            # 최악이다. known_* 는 이번 실행에 실제로 존재한 행·컬럼이다.
            self.override_conflicts = applier.conflicts(
                known_row_ids={make_row_id(i) for i in df.index},
                known_cols=set(cleaned_col_vals.keys()),
            )
            print(f"손 편집 반영: {applier.applied_count}건")
            if self.override_conflicts:
                print(f"[경고] 적용하지 못한 편집 {len(self.override_conflicts)}건 — 검토 필요")

        # ── dashboard.json 차트 → summary.sections 자동 생성 ─────────────────
        # 웹 관리자 화면(Step2)은 dashboard.json의 charts만 편집할 뿐 config.yaml의
        # summary.sections를 편집하는 UI가 없어, 웹으로 만든 프로젝트는 sections가
        # 항상 비어 있다. sections가 비어 있으면 아래 SummarySheetWriter 자체가
        # 건너뛰어져 Summary 시트가 생성되지 않고, 결과적으로 대시보드에 설정한
        # 차트가 엑셀에는 하나도 반영되지 않는다(에러 없이 조용히 스킵됨).
        # 사용자가 summary.sections를 직접 구성하지 않은 경우에 한해, dashboard.json의
        # 차트 목록에서 필요한 unique_count 섹션을 자동으로 만들어 붙여준다.
        if not cfg.summary.sections and proj_dir:
            dashboard_path = proj_dir / "dashboard.json"
            if dashboard_path.exists():
                try:
                    with open(dashboard_path, encoding="utf-8") as f:
                        dash_data = json.load(f)
                    from engine.config import SummarySection, SummaryLayout, SummaryColumnItem
                    chart_sections = []
                    seen: set[str] = set()
                    
                    for c in dash_data.get("charts", []):
                        c_type = c.get("type")
                        if c_type == "multibar":
                            cols = c.get("cols", [])
                            valid_items = []
                            for ci in cols:
                                col_name = ci.get("col")
                                if col_name and col_name in col_index_map:
                                    valid_items.append(
                                        SummaryColumnItem(col_ref=col_name, label=ci.get("label") or col_name)
                                    )
                            if valid_items:
                                chart_sections.append(
                                    SummarySection(
                                        title=c.get("title") or "다중 항목 집계",
                                        type="binary_sum",
                                        columns=valid_items,
                                    )
                                )
                            continue
                            
                        col_ref = c.get("col") or c.get("colRef")
                        if not col_ref or col_ref in seen or col_ref not in col_index_map:
                            continue
                        seen.add(col_ref)
                        chart_sections.append(
                            SummarySection(
                                title=c.get("title") or f"{col_ref} 집계",
                                type="unique_count",
                                col_ref=col_ref,
                            )
                        )

                    if chart_sections:
                        if cfg.summary.layout is None:
                            cfg.summary.layout = SummaryLayout()
                        n_layout_cols = max(cfg.summary.layout.cols, 1)
                        for i, sec in enumerate(chart_sections):
                            sec.layout_col = (i % n_layout_cols) + 1
                        cfg.summary.sections = chart_sections
                        print(
                            f"[정보] dashboard.json 차트 {len(chart_sections)}개 기준으로 "
                            f"summary.sections 자동 생성"
                        )
                except Exception as exc:
                    print(f"[경고] dashboard.json 기반 summary.sections 자동 생성 실패: {exc}")

        if cfg.summary.sections:
            SummarySheetWriter(cfg).write(out_wb, df, col_index_map, cleaned_col_vals)

        # ── 원본 시트 ─────────────────────────────────────────────────────────
        _write_raw_sheet(out_wb, src_wb, cfg, raw_df, source_headers=source_headers)

        # ── Config / Guide sheets ─────────────────────────────────────────────
        from engine.config_excel import write_config_sheet, write_guide_sheet
        write_config_sheet(out_wb, cfg)
        write_guide_sheet(out_wb)

        # ── output path ───────────────────────────────────────────────────────
        out_dir  = Path(cfg.paths.output_dir)
        if not out_dir.is_absolute():
            if self._config_path:
                out_dir = self._config_path.parent / out_dir
            else:
                # config_path가 없는 경우 (CLI Excel 다이렉트 호출 등)
                # 서브프로젝트 폴더가 존재하면 그 하위 output으로 똑똑하게 스마트 폴백
                proj_folder = Path("storage/projects") / cfg.project
                if proj_folder.exists():
                    out_dir = proj_folder / out_dir
                else:
                    out_dir = Path(cfg.paths.output_dir)
        
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / cfg.paths.output_file

        if dry_run:
            print("\n[dry-run] 파일 저장을 건너뛰었습니다. (정상 작동 확인)")
            print(f"완료: {n_rows}건 (dry-run)")
            return None

        try:
            out_wb.save(out_path)
            save_path = out_path
        except PermissionError:
            ts = datetime.now().strftime("%H%M%S")
            save_path = out_path.with_stem(out_path.stem + f"_{ts}")
            out_wb.save(save_path)
            print(f"[주의] 파일이 열려 있어 임시 저장: {save_path.name}")

        print(f"저장: {save_path}")

        # ── slicers ───────────────────────────────────────────────────────────
        include_slicers = getattr(cfg.excel_options, "include_slicers", True)
        if include_slicers:
            if not cfg.slicers:
                from engine.config import SlicerDef
                cfg.slicers = [SlicerDef(col=cd.output_col) for cd in cfg.columns if cd.include_in_slicer]

            if cfg.slicers:
                inject_slicers(save_path, cfg, col_index_map)
        else:
            print("[정보] 엑셀 슬라이서 포함 설정이 비활성화되어 슬라이서 주입을 건너뜁니다.")

        # ── config.json 저장 (프로젝트 폴더) ─────────────────────────────────
        if self._config_path:
            config_json = self._config_path.parent / f"{cfg.project}_config.json"
            try:
                config_json.parent.mkdir(parents=True, exist_ok=True)
                with open(config_json, "w", encoding="utf-8") as f:
                    json.dump(
                        cfg.model_dump(exclude_none=True),
                        f, ensure_ascii=False, indent=2, default=str,
                    )
                print(f"설정 JSON: {config_json}")
            except Exception as exc:
                print(f"[주의] config.json 저장 실패: {exc}")

        print(f"\n완료: {n_rows}건 → {save_path.name}")
        return save_path
