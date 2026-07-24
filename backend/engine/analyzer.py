from __future__ import annotations

from typing import Any

import openpyxl

from engine.base_analyzer import BaseAnalyzer


_MAX_HEADER_SEARCH = 5   # look up to this many rows for the header
_DATA_SCAN_ROWS    = 5   # rows to scan after header for actual data start
_DATA_MIN_DENSITY  = 0.1 # minimum fill ratio to consider a row as real data


def _row_cells(ws, row: int) -> list[Any]:
    return [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]


def _score_row(cells: list[Any]) -> float:
    """Score a row as a candidate header (0.0–1.0)."""
    if not cells:
        return 0.0
    total = len(cells)
    non_empty = sum(1 for c in cells if c is not None and str(c).strip() != "")
    strings   = sum(1 for c in cells if isinstance(c, str) and c.strip() != "")
    numerics  = sum(1 for c in cells if isinstance(c, (int, float)))

    ne_ratio  = non_empty / total
    str_ratio = strings / total if non_empty else 0.0
    no_num    = 0.2 if numerics == 0 else 0.0

    return round(ne_ratio * 0.4 + str_ratio * 0.4 + no_num, 4)


def _detect_data_start(ws, header_row: int) -> int:
    """Find first row after header that looks like real data.

    Strategy:
    1. Collect fill densities for the next _DATA_SCAN_ROWS rows.
    2. Use max_density * 0.5 as adaptive threshold (handles both dense
       and sparse surveys without a fixed magic number).
    3. Fallback: header_row + 1 when all scanned rows are empty.
    """
    total = ws.max_column or 1
    densities: list[tuple[int, float]] = []
    for r in range(header_row + 1, header_row + _DATA_SCAN_ROWS + 1):
        cells = _row_cells(ws, r)
        non_empty = sum(1 for c in cells if c is not None and str(c).strip() != "")
        densities.append((r, non_empty / total))

    if not densities:
        return header_row + 1

    max_density = max(d for _, d in densities)
    if max_density < _DATA_MIN_DENSITY:
        return header_row + 1  # all rows are nearly empty — give up

    threshold = max_density * 0.5
    for r, density in densities:
        if density >= threshold:
            return r
    return header_row + 1


class ExcelAnalyzer(BaseAnalyzer):
    """Inspect an xlsx file and propose source configuration.

    병합 셀·다중 타이틀 행처럼 사람이 손으로 만든 xlsx의 "머리말이 몇 번째 줄인지 애매한"
    문제를 점수 매겨 추정한다(`_score_row`/`_detect_data_start`) — 이 모호함 자체가
    엑셀에만 있는 문제라, 이 클래스에만 존재하는 로직이다. `CsvAnalyzer`는 이 모호함이
    구조적으로 없어 헤더를 1행 고정으로 둔다(csv_analyzer.py 참고).
    """

    def sheet_names(self) -> list[str]:
        wb = openpyxl.load_workbook(self._path, read_only=True, data_only=True)
        names = wb.sheetnames
        wb.close()
        return names

    def detect_header_row(self, sheet_name: str | None = None) -> dict:
        """Return detection result dict with header_row, scores, sample."""
        wb = openpyxl.load_workbook(self._path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        if ws is None:
            wb.close()
            return {}

        scores: list[tuple[int, float, list]] = []
        for r in range(1, _MAX_HEADER_SEARCH + 1):
            cells = _row_cells(ws, r)
            scores.append((r, _score_row(cells), cells))

        best_row, best_score, best_cells = max(scores, key=lambda x: x[1])
        data_start = _detect_data_start(ws, best_row)

        wb.close()
        return {
            "sheet": ws.title,
            "header_row": best_row,
            "header_score": best_score,
            "data_start_row": data_start,
            "column_count": len([c for c in best_cells if c is not None]),
            "sample_headers": [str(c) for c in best_cells if c is not None][:10],
            "all_scores": [(r, sc) for r, sc, _ in scores],
        }

    def all_headers(self, sheet_name: str | None = None) -> list[str | None]:
        """Return every cell value from the detected header row (None for empty)."""
        wb = openpyxl.load_workbook(self._path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        if ws is None:
            wb.close()
            return []
        detection = self.detect_header_row(sheet_name)
        hrow = detection.get("header_row", 1)
        cells = _row_cells(ws, hrow)
        wb.close()
        return cells
