import yaml
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import TableStyleInfo

_DEFAULT_STYLE: dict = {
    "font": {"name": "Arial"},
    "cleaned": {
        "subtotal_row": {
            "height": 20,
            "bg_color": "E2EFDA",
            "font_color": "375623",
            "font_size": 9,
            "bold": True,
            "border_side_color": "93C47D",
            "border_top_bottom_color": "6AA84F",
        },
        "header_row": {
            "height": 28,
            "bg_color": "1F4E79",
            "font_color": "FFFFFF",
            "font_size": 10,
            "bold": True,
        },
        "data_row": {"font_size": 9},
        "table": {
            "style_name": "TableStyleMedium2",
            "show_row_stripes": True,
            "show_column_stripes": False,
        },
        "freeze_pane": "A3",
    },
    "summary": {
        "section_title": {
            "bg_color": "2E75B6",
            "font_color": "FFFFFF",
            "font_size": 11,
            "bold": True,
        },
        "header": {
            "bg_color": "BDD7EE",
            "font_color": "1F4E79",
            "font_size": 10,
            "bold": True,
        },
        "data": {"font_size": 10},
        "cell_border_color": "CCCCCC",
        "column_widths": {
            "A": 26, "B": 16, "C": 10, "D": 2,
            "E": 22, "F": 16, "G": 10, "H": 12,
        },
        "freeze_pane": "A1",
    },
}


def load_style(config_path: str | Path | None, style_file: str | None) -> dict:
    if not style_file:
        return _DEFAULT_STYLE
    p = Path(style_file)
    if not p.is_absolute() and config_path:
        p = Path(config_path).parent / style_file
    if not p.exists():
        return _DEFAULT_STYLE
    with open(p, encoding="utf-8") as f:
        return yaml.safe_load(f) or _DEFAULT_STYLE


def build_cleaned_styles(st: dict) -> dict:
    fn  = st.get("font", {}).get("name", "Arial")
    cs  = st.get("cleaned", {})
    sub = cs.get("subtotal_row", {})
    hdr = cs.get("header_row", {})
    dat = cs.get("data_row", {})
    tbl = cs.get("table", {})
    sc  = sub.get("border_side_color", "93C47D")
    tc  = sub.get("border_top_bottom_color", "6AA84F")
    return {
        "sub_fill":   PatternFill("solid", start_color=sub.get("bg_color", "E2EFDA")),
        "sub_font":   Font(bold=sub.get("bold", True), name=fn,
                          size=sub.get("font_size", 9),
                          color=sub.get("font_color", "375623")),
        "sub_align":  Alignment(horizontal="center", vertical="center"),
        "sub_bdr":    Border(
            left=Side(style="thin",   color=sc),
            right=Side(style="thin",  color=sc),
            top=Side(style="medium",  color=tc),
            bottom=Side(style="medium", color=tc),
        ),
        "sub_height": sub.get("height", 20),
        "hdr_fill":   PatternFill("solid", start_color=hdr.get("bg_color", "1F4E79")),
        "hdr_font":   Font(bold=hdr.get("bold", True), name=fn,
                          size=hdr.get("font_size", 10),
                          color=hdr.get("font_color", "FFFFFF")),
        "hdr_align":  Alignment(horizontal="center", vertical="center", wrap_text=True),
        "hdr_height": hdr.get("height", 28),
        "dat_font":   Font(name=fn, size=dat.get("font_size", 9)),
        "dat_align":  Alignment(vertical="center"),
        "tbl_style":  TableStyleInfo(
            name=tbl.get("style_name", "TableStyleMedium2"),
            showFirstColumn=False, showLastColumn=False,
            showRowStripes=tbl.get("show_row_stripes", True),
            showColumnStripes=tbl.get("show_column_stripes", False),
        ),
        "freeze_pane": cs.get("freeze_pane", "A3"),
    }


def build_summary_styles(st: dict) -> dict:
    fn  = st.get("font", {}).get("name", "Arial")
    ss  = st.get("summary", {})
    sec = ss.get("section_title", {})
    hdr = ss.get("header", {})
    dat = ss.get("data", {})
    bc  = ss.get("cell_border_color", "CCCCCC")
    thin = Side(style="thin", color=bc)
    return {
        "sec_fill":   PatternFill("solid", start_color=sec.get("bg_color", "2E75B6")),
        "sec_font":   Font(bold=sec.get("bold", True), name=fn,
                          size=sec.get("font_size", 11),
                          color=sec.get("font_color", "FFFFFF")),
        "hdr_fill":   PatternFill("solid", start_color=hdr.get("bg_color", "BDD7EE")),
        "hdr_font":   Font(bold=hdr.get("bold", True), name=fn,
                          size=hdr.get("font_size", 10),
                          color=hdr.get("font_color", "1F4E79")),
        "dat_font":   Font(name=fn, size=dat.get("font_size", 10)),
        "ctr":        Alignment(horizontal="center", vertical="center"),
        "lft":        Alignment(horizontal="left",   vertical="center"),
        "bdr":        Border(left=thin, right=thin, top=thin, bottom=thin),
        "col_widths": ss.get("column_widths", {
            "A": 26, "B": 16, "C": 10, "D": 2, "E": 22, "F": 16, "G": 10, "H": 12,
        }),
        "freeze_pane": ss.get("freeze_pane", "A1"),
    }
