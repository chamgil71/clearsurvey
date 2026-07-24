"""
CSV 업로드 지원 테스트.

Coverage targets:
  - CsvAnalyzer                — 헤더 탐지·인코딩 감지
  - engine.pipeline._csv_to_dataframe / _coerce_csv_value
  - SurveyPipeline.run() with csv 입력 — xlsx 입력과 결과가 동일해야 한다

실행: pytest tests/test_config_csv.py -v
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from engine.config import SurveyConfig
from engine.csv_analyzer import CsvAnalyzer, detect_csv_encoding
from engine.pipeline import SurveyPipeline, _coerce_csv_value


# ─────────────────────────────────────────────────────────────────────────────
# 헬퍼
# ─────────────────────────────────────────────────────────────────────────────

_ROWS = [
    ("이름", "부서", "점수", "연락처"),
    ("홍길동", "개발팀", "90", "01012345678"),
    ("김철수", "영업팀", "85", "01098765432"),
    ("이영희", "개발팀", "92", "01055512345"),
]


def _write_csv(path: Path, encoding: str = "utf-8-sig") -> Path:
    with open(path, "w", encoding=encoding, newline="") as f:
        for row in _ROWS:
            f.write(",".join(row) + "\r\n")
    return path


def _cfg_for(data_path: Path, tmp_path: Path) -> SurveyConfig:
    return SurveyConfig(
        project="csv_test",
        source={"sheet": None, "header_row": 1, "data_start_row": 2, "file": str(data_path)},
        paths={"output_dir": str(tmp_path / "output"), "output_file": "result.xlsx"},
        sheets={"cleaned": "Cleaned", "summary": "Summary"},
        columns=[
            {"output_col": "이름", "source_col": 1, "transform": "copy"},
            {"output_col": "부서", "source_col": 2, "transform": "copy"},
            {"output_col": "점수", "source_col": 3, "transform": "copy"},
            {"output_col": "연락처", "source_col": 4, "transform": "copy"},
        ],
    )


# ─────────────────────────────────────────────────────────────────────────────
# CsvAnalyzer
# ─────────────────────────────────────────────────────────────────────────────

class TestCsvAnalyzer:
    def test_sheet_names_returns_virtual_sheet(self, tmp_path):
        path = _write_csv(tmp_path / "data.csv")
        analyzer = CsvAnalyzer(path)
        assert analyzer.sheet_names() == ["CSV"]

    def test_detect_header_row_is_always_row_1(self, tmp_path):
        path = _write_csv(tmp_path / "data.csv")
        analyzer = CsvAnalyzer(path)
        detection = analyzer.detect_header_row()
        assert detection["header_row"] == 1
        assert detection["data_start_row"] == 2
        assert detection["column_count"] == 4

    def test_all_headers_matches_first_row(self, tmp_path):
        path = _write_csv(tmp_path / "data.csv")
        analyzer = CsvAnalyzer(path)
        assert analyzer.all_headers() == ["이름", "부서", "점수", "연락처"]

    def test_generate_config_yaml_suggests_phone_transform(self, tmp_path):
        """'연락처' 헤더 → normalize_phone 자동 제안(base_analyzer의 _suggest_transform 공유 확인)."""
        import yaml
        path = _write_csv(tmp_path / "data.csv")
        analyzer = CsvAnalyzer(path)
        out = analyzer.generate_config_yaml(tmp_path / "config.yaml", project_name="p")
        cfg = yaml.safe_load(out.read_text(encoding="utf-8"))
        phone_col = next(c for c in cfg["columns"] if c["output_col"] == "연락처")
        assert phone_col["transform"] == "normalize_phone"

    def test_encoding_detection_cp949(self, tmp_path):
        path = _write_csv(tmp_path / "data_cp949.csv", encoding="cp949")
        enc = detect_csv_encoding(path)
        assert enc in ("cp949", "euc-kr")  # 둘 다 이 데이터를 올바르게 디코딩한다
        analyzer = CsvAnalyzer(path)
        assert analyzer.all_headers() == ["이름", "부서", "점수", "연락처"]


# ─────────────────────────────────────────────────────────────────────────────
# _coerce_csv_value — 앞자리 0 보존
# ─────────────────────────────────────────────────────────────────────────────

class TestCoerceCsvValue:
    def test_plain_integer_becomes_int(self):
        assert _coerce_csv_value("90") == 90
        assert isinstance(_coerce_csv_value("90"), int)

    def test_decimal_becomes_float(self):
        assert _coerce_csv_value("3.14") == 3.14

    def test_leading_zero_phone_stays_string(self):
        """01012345678처럼 앞자리 0이 있는 숫자열은 정수로 바뀌면 안 된다(0 소실 방지)."""
        assert _coerce_csv_value("01012345678") == "01012345678"
        assert isinstance(_coerce_csv_value("01012345678"), str)

    def test_zero_alone_becomes_int(self):
        assert _coerce_csv_value("0") == 0

    def test_empty_string_becomes_none(self):
        assert _coerce_csv_value("") is None
        assert _coerce_csv_value("   ") is None

    def test_negative_integer(self):
        assert _coerce_csv_value("-42") == -42

    def test_non_numeric_stays_string(self):
        assert _coerce_csv_value("개발팀") == "개발팀"


# ─────────────────────────────────────────────────────────────────────────────
# SurveyPipeline.run() with csv — xlsx 결과와 동등해야 한다
# ─────────────────────────────────────────────────────────────────────────────

class TestPipelineCsvVsXlsxEquivalence:
    def test_csv_pipeline_produces_same_cleaned_rows_as_xlsx(self, tmp_path):
        """같은 데이터를 xlsx/csv로 각각 넣었을 때 Cleaned 시트 값이 동일해야 한다.

        docs/plan/pending/new_beginnings_comparison_plan.md §1-B가 주장하는
        "_sheet_to_dataframe 계약만 지키면 이후 로직은 완전히 동일하다"를 직접 검증한다.
        """
        # xlsx 소스
        xlsx_path = tmp_path / "data.xlsx"
        wb = Workbook()
        ws = wb.active
        for ci, header in enumerate(_ROWS[0], 1):
            ws.cell(1, ci, header)
        for ri, row in enumerate(_ROWS[1:], 2):
            ws.cell(ri, 1, row[0])
            ws.cell(ri, 2, row[1])
            ws.cell(ri, 3, int(row[2]))       # openpyxl: 점수는 원래 타입 있는 int로 온다
            ws.cell(ri, 4, row[3])            # 연락처는 텍스트로 저장(엑셀에서도 흔한 관례)
        wb.save(xlsx_path)

        # csv 소스 (전부 문자열)
        csv_path = _write_csv(tmp_path / "data.csv")

        cfg_xlsx = _cfg_for(xlsx_path, tmp_path / "xlsx_out")
        cfg_csv  = _cfg_for(csv_path, tmp_path / "csv_out")

        out_xlsx = SurveyPipeline(cfg_xlsx).run(input_path=xlsx_path)
        out_csv  = SurveyPipeline(cfg_csv).run(input_path=csv_path)

        def _cleaned_rows(path: Path) -> list[tuple]:
            wb = openpyxl.load_workbook(path)
            ws = wb["Cleaned"]
            return [
                r for r in ws.iter_rows(min_row=3, max_col=4, values_only=True)
                if any(c is not None for c in r)
            ]

        rows_xlsx = _cleaned_rows(out_xlsx)
        rows_csv  = _cleaned_rows(out_csv)
        assert rows_xlsx == rows_csv

    def test_csv_phone_column_keeps_leading_zero_through_pipeline(self, tmp_path):
        csv_path = _write_csv(tmp_path / "data.csv")
        cfg = _cfg_for(csv_path, tmp_path)
        out = SurveyPipeline(cfg).run(input_path=csv_path)
        wb = openpyxl.load_workbook(out)
        ws = wb["Cleaned"]
        phones = [
            r[3] for r in ws.iter_rows(min_row=3, max_col=4, values_only=True)
            if r[3] is not None
        ]
        assert "01012345678" in phones

    def test_csv_numeric_column_is_numeric_type(self, tmp_path):
        """점수 컬럼이 문자열 "text"가 아니라 실제 숫자로 들어가야 KPI 합산·서식이 정상 동작한다."""
        csv_path = _write_csv(tmp_path / "data.csv")
        cfg = _cfg_for(csv_path, tmp_path)
        out = SurveyPipeline(cfg).run(input_path=csv_path)
        wb = openpyxl.load_workbook(out)
        ws = wb["Cleaned"]
        scores = [
            r[2] for r in ws.iter_rows(min_row=3, max_col=4, values_only=True)
            if r[2] is not None
        ]
        assert all(isinstance(s, (int, float)) for s in scores)
        assert 90 in scores
