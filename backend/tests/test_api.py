"""
Tests for app/main.py — FastAPI endpoints

Coverage targets:
  - GET  /api/health                   — 생존 확인
  - GET  /api/projects/{name}/status   — idle 기본값, 유효성 검증
  - GET  /api/projects/{name}/config   — 404 (미존재), config 로드
  - POST /api/projects/{name}/config   — 저장 성공, 400 검증 실패, 404 미존재
  - POST /api/projects/{name}/run      — 404 프로젝트 없음, 404 config 없음, 409 중복실행, started
  - Security: path traversal → 400

실행: pytest tests/test_api.py -v
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
import yaml
from fastapi.testclient import TestClient
from openpyxl import Workbook

# app/main.py를 import하기 전에 sys.path에 프로젝트 루트 추가
import sys
_PROJECT_ROOT = Path(__file__).parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))
    sys.path.insert(0, str(_PROJECT_ROOT / "app"))

from app.main import app, _job_set, _pipeline_jobs, _build_single_html_template


# ─────────────────────────────────────────────────────────────────────────────
# 픽스처
# ─────────────────────────────────────────────────────────────────────────────

@pytest.fixture
def client(tmp_path, monkeypatch):
    """TestClient with roots monkeypatched to tmp_path."""
    import app.main as main_module
    monkeypatch.setattr(main_module, "STORAGE_ROOT", tmp_path / "storage")
    monkeypatch.setattr(main_module, "FRONTEND_ROOT", tmp_path / "frontend")
    monkeypatch.setattr(main_module, "BACKEND_ROOT", tmp_path / "backend")
    # 잡 상태 초기화
    _pipeline_jobs.clear()
    with TestClient(app) as c:
        yield c


@pytest.fixture
def project_dir(tmp_path):
    """tmp_path 기반 프로젝트 디렉터리를 반환합니다."""
    d = tmp_path / "storage" / "projects" / "demo"
    d.mkdir(parents=True, exist_ok=True)
    (d / "output").mkdir(exist_ok=True)
    return d


def _write_valid_config(proj_dir: Path, data_xlsx: Path | None = None) -> Path:
    """proj_dir/config.yaml 을 작성하고 경로를 반환합니다."""
    cfg: dict = {
        "project": proj_dir.name,
        "source": {
            "file": str(data_xlsx) if data_xlsx else None,
            "sheet": None,
            "header_row": 1,
            "data_start_row": 2,
        },
        "paths": {"output_dir": "output", "output_file": "result.xlsx"},
        "sheets": {"cleaned": "Cleaned", "summary": "Summary"},
        "columns": [
            {"output_col": "이름", "source_col": 1, "transform": "copy"},
        ],
    }
    config_yaml = proj_dir / "config.yaml"
    config_yaml.write_text(yaml.dump(cfg, allow_unicode=True), encoding="utf-8")
    return config_yaml


def _make_xlsx(path: Path) -> None:
    """단순 2행 xlsx 파일을 생성합니다."""
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1, "이름")
    ws.cell(2, 1, "홍길동")
    wb.save(path)


# ─────────────────────────────────────────────────────────────────────────────
# GET /api/health
# ─────────────────────────────────────────────────────────────────────────────

class TestHealthCheck:
    def test_returns_200(self, client):
        resp = client.get("/api/health")
        assert resp.status_code == 200

    def test_status_is_ok(self, client):
        resp = client.get("/api/health")
        assert resp.json()["status"] == "ok"

    def test_has_time_field(self, client):
        resp = client.get("/api/health")
        assert "time" in resp.json()


# ─────────────────────────────────────────────────────────────────────────────
# GET /api/projects/{name}/status
# ─────────────────────────────────────────────────────────────────────────────

class TestGetPipelineStatus:
    def test_unknown_project_returns_idle(self, client):
        resp = client.get("/api/projects/unknown_proj/status")
        assert resp.status_code == 200
        assert resp.json()["status"] == "idle"

    def test_invalid_name_returns_400(self, client):
        resp = client.get("/api/projects/../etc/status")
        assert resp.status_code in (400, 404)

    def test_special_chars_returns_400(self, client):
        resp = client.get("/api/projects/bad%00name/status")
        assert resp.status_code in (400, 422)

    def test_reflects_set_status(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        _job_set("my_proj", {"status": "running", "started_at": "2026-01-01T00:00:00"})
        resp = client.get("/api/projects/my_proj/status")
        assert resp.json()["status"] == "running"
        _pipeline_jobs.clear()


# ─────────────────────────────────────────────────────────────────────────────
# GET /api/projects/{name}/freshness
# ─────────────────────────────────────────────────────────────────────────────

class TestGetProjectFreshness:
    def test_unknown_project_returns_404(self, client):
        resp = client.get("/api/projects/unknown_proj/freshness")
        assert resp.status_code == 404

    def test_invalid_name_returns_400(self, client):
        resp = client.get("/api/projects/../etc/freshness")
        assert resp.status_code in (400, 404)

    def test_no_output_yet_is_not_stale(self, client, tmp_path):
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        (proj_dir / "output").mkdir()
        _write_valid_config(proj_dir)

        resp = client.get("/api/projects/demo/freshness")
        assert resp.status_code == 200
        body = resp.json()
        assert body["has_output"] is False
        assert body["is_stale"] is False
        assert body["output_generated_at"] is None
        assert body["config_updated_at"] is not None

    def test_stale_when_config_newer_than_output(self, client, tmp_path):
        import os

        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        (proj_dir / "output").mkdir()
        config_path = _write_valid_config(proj_dir)
        output_path = proj_dir / "output" / "result.xlsx"
        output_path.write_text("dummy")

        now = 1_800_000_000
        os.utime(output_path, (now - 100, now - 100))
        os.utime(config_path, (now, now))

        resp = client.get("/api/projects/demo/freshness")
        body = resp.json()
        assert body["has_output"] is True
        assert body["is_stale"] is True

    def test_not_stale_when_output_newer_than_config(self, client, tmp_path):
        import os

        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        (proj_dir / "output").mkdir()
        config_path = _write_valid_config(proj_dir)
        output_path = proj_dir / "output" / "result.xlsx"
        output_path.write_text("dummy")

        now = 1_800_000_000
        os.utime(config_path, (now - 100, now - 100))
        os.utime(output_path, (now, now))

        resp = client.get("/api/projects/demo/freshness")
        body = resp.json()
        assert body["has_output"] is True
        assert body["is_stale"] is False

    def test_stale_when_dashboard_newer_than_output(self, client, tmp_path):
        import os

        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        (proj_dir / "output").mkdir()
        config_path = _write_valid_config(proj_dir)
        output_path = proj_dir / "output" / "result.xlsx"
        output_path.write_text("dummy")
        dashboard_path = proj_dir / "dashboard.json"
        dashboard_path.write_text("{}")

        now = 1_800_000_000
        os.utime(config_path, (now - 100, now - 100))
        os.utime(output_path, (now - 100, now - 100))
        os.utime(dashboard_path, (now, now))

        resp = client.get("/api/projects/demo/freshness")
        body = resp.json()
        assert body["is_stale"] is True


# ─────────────────────────────────────────────────────────────────────────────
# GET /api/projects/{name}/config
# ─────────────────────────────────────────────────────────────────────────────

class TestGetProjectConfig:
    def test_nonexistent_project_returns_404(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        (tmp_path / "storage" / "projects").mkdir(parents=True, exist_ok=True)
        resp = client.get("/api/projects/no_such_proj/config")
        assert resp.status_code == 404

    def test_invalid_name_returns_400(self, client):
        resp = client.get("/api/projects/../config")
        assert resp.status_code in (400, 404)

    def test_valid_project_returns_config(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        _write_valid_config(proj_dir)
        resp = client.get("/api/projects/demo/config")
        assert resp.status_code == 200
        body = resp.json()
        assert "config" in body
        assert "dashboard" in body

    def test_config_contains_project_name(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        _write_valid_config(proj_dir)
        resp = client.get("/api/projects/demo/config")
        assert resp.json()["config"]["project"] == "demo"

    def test_dashboard_json_included_when_present(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        _write_valid_config(proj_dir)
        custom_dash = {"version": 99, "kpi": [], "charts": [], "list": {"visible_cols": [], "filter_cols": []}}
        (proj_dir / "dashboard.json").write_text(json.dumps(custom_dash), encoding="utf-8")
        resp = client.get("/api/projects/demo/config")
        assert resp.json()["dashboard"]["version"] == 99


# ─────────────────────────────────────────────────────────────────────────────
# POST /api/projects/{name}/config — 설정 저장
# ─────────────────────────────────────────────────────────────────────────────

class TestSaveProjectConfig:
    def _valid_payload(self) -> dict:
        return {
            "config": {
                "project": "demo",
                "source": {"header_row": 1, "data_start_row": 2},
                "paths": {"output_dir": "output", "output_file": "result.xlsx"},
                "sheets": {"cleaned": "Cleaned", "summary": "Summary"},
                "columns": [{"output_col": "이름", "source_col": 1, "transform": "copy"}],
            },
            "dashboard": None,
        }

    def test_nonexistent_project_returns_404(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        (tmp_path / "storage" / "projects").mkdir(parents=True, exist_ok=True)
        resp = client.post("/api/projects/no_such/config", json=self._valid_payload())
        assert resp.status_code == 404

    def test_invalid_project_name_returns_400(self, client):
        resp = client.post("/api/projects/../../etc/config", json=self._valid_payload())
        assert resp.status_code in (400, 404)

    def test_valid_save_returns_success(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        resp = client.post("/api/projects/demo/config", json=self._valid_payload())
        assert resp.status_code == 200
        assert resp.json()["status"] == "success"

    def test_config_yaml_written_on_success(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        client.post("/api/projects/demo/config", json=self._valid_payload())
        assert (proj_dir / "config.yaml").exists()

    def test_empty_config_returns_400(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        resp = client.post("/api/projects/demo/config", json={"config": None, "dashboard": None})
        assert resp.status_code == 400

    def test_dashboard_json_written_when_provided(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        payload = self._valid_payload()
        payload["dashboard"] = {
            "version": 42, "kpi": [], "charts": [],
            "list": {"visible_cols": [], "filter_cols": []}
        }
        client.post("/api/projects/demo/config", json=payload)
        assert (proj_dir / "dashboard.json").exists()
        data = json.loads((proj_dir / "dashboard.json").read_text(encoding="utf-8"))
        assert data["version"] == 42


# ─────────────────────────────────────────────────────────────────────────────
# POST /api/projects/{name}/run — 파이프라인 실행
# ─────────────────────────────────────────────────────────────────────────────

class TestRunProjectPipeline:
    def test_nonexistent_project_returns_404(self, client, tmp_path, monkeypatch):
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        (tmp_path / "storage" / "projects").mkdir(parents=True, exist_ok=True)
        resp = client.post("/api/projects/ghost/run")
        assert resp.status_code == 404

    def test_invalid_name_returns_400(self, client):
        resp = client.post("/api/projects/../run")
        assert resp.status_code in (400, 404)

    def test_missing_config_yaml_returns_404(self, client, tmp_path, monkeypatch):
        """프로젝트 폴더는 있지만 config.yaml 없으면 404."""
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "nocfg"
        proj_dir.mkdir(parents=True)
        resp = client.post("/api/projects/nocfg/run")
        assert resp.status_code == 404

    def test_returns_started_status(self, client, tmp_path, monkeypatch):
        """config.yaml 이 있으면 {status: 'started'} 를 즉시 반환한다."""
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        data_xlsx = tmp_path / "data.xlsx"
        _make_xlsx(data_xlsx)
        _write_valid_config(proj_dir, data_xlsx=data_xlsx)
        resp = client.post("/api/projects/demo/run")
        assert resp.status_code == 200
        assert resp.json()["status"] == "started"

    def test_conflict_409_when_already_running(self, client, tmp_path, monkeypatch):
        """이미 실행 중이면 409 Conflict 반환."""
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)
        data_xlsx = tmp_path / "data.xlsx"
        _make_xlsx(data_xlsx)
        _write_valid_config(proj_dir, data_xlsx=data_xlsx)
        # 잡을 직접 'running' 상태로 설정
        _job_set("demo", {"status": "running", "started_at": "2026-01-01T00:00:00"})
        resp = client.post("/api/projects/demo/run")
        assert resp.status_code == 409
        _pipeline_jobs.clear()

    def test_broken_config_returns_400(self, client, tmp_path, monkeypatch):
        """Pydantic 검증 실패 config.yaml 이면 400을 반환한다.

        source_col/source_cols/source_label 없는 컬럼은 ColumnDef._check_source
        validator가 ValidationError 를 발생시키므로 반드시 400을 반환해야 한다.
        """
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        proj_dir = tmp_path / "storage" / "projects" / "broken"
        proj_dir.mkdir(parents=True)
        # source 정보 없는 컬럼 → ColumnDef._check_source → ValidationError → 400
        broken_yaml = (
            "project: broken\n"
            "columns:\n"
            "  - output_col: bad_col\n"
            "    # source_col / source_cols / source_label 모두 없음 → Pydantic 오류\n"
        )
        (proj_dir / "config.yaml").write_text(broken_yaml, encoding="utf-8")
        resp = client.post("/api/projects/broken/run")
        assert resp.status_code == 400


# ─────────────────────────────────────────────────────────────────────────────
# Security: 프로젝트 이름 경로 순회 방지
# ─────────────────────────────────────────────────────────────────────────────

class TestProjectNameSecurity:
    PATH_TRAVERSAL_NAMES = [
        "../etc",
        "../../windows/system32",
        "a" * 65,           # 64자 초과
        "hello world",      # 공백 포함
        "proj;drop",        # 세미콜론
    ]

    @pytest.mark.parametrize("bad_name", PATH_TRAVERSAL_NAMES)
    def test_status_rejects_unsafe_names(self, client, bad_name):
        encoded = bad_name.replace("/", "%2F").replace(" ", "%20")
        resp = client.get(f"/api/projects/{encoded}/status")
        assert resp.status_code in (400, 404, 422)

    @pytest.mark.parametrize("bad_name", PATH_TRAVERSAL_NAMES)
    def test_run_rejects_unsafe_names(self, client, bad_name):
        encoded = bad_name.replace("/", "%2F").replace(" ", "%20")
        resp = client.post(f"/api/projects/{encoded}/run")
        assert resp.status_code in (400, 404, 422)

    def test_valid_korean_name_passes_validation(self, client, tmp_path, monkeypatch):
        """한글 프로젝트명은 허용된다."""
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        (tmp_path / "storage" / "projects").mkdir(parents=True, exist_ok=True)
        resp = client.get("/api/projects/테스트프로젝트/status")
        # 200 idle 또는 404 (프로젝트 없음)가 모두 허용됨; 400이면 안 됨
        assert resp.status_code != 400


# ─────────────────────────────────────────────────────────────────────────────
# SSE 로그 스트림 엔드포인트
# ─────────────────────────────────────────────────────────────────────────────

class TestStreamProjectLogs:
    """GET /api/projects/{name}/logs/stream — SSE 실시간 로그 스트리밍."""

    def test_invalid_name_returns_400(self, client):
        # Names with illegal characters must be rejected before streaming
        resp = client.get("/api/projects/hello%20world/logs/stream")
        assert resp.status_code == 400

    def test_returns_event_stream_content_type(self, client, tmp_path, monkeypatch):
        """SSE 엔드포인트는 text/event-stream 미디어 타입을 반환한다."""
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        _pipeline_jobs.clear()
        _job_set("demo", {"status": "done"})

        with client.stream("GET", "/api/projects/demo/logs/stream") as resp:
            assert resp.status_code == 200
            assert "text/event-stream" in resp.headers.get("content-type", "")

    def test_done_job_terminates_stream(self, client, tmp_path, monkeypatch):
        """완료된 잡은 스트림에 종료 메시지를 보내고 종료된다."""
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        _pipeline_jobs.clear()
        _job_set("demo", {"status": "done"})

        with client.stream("GET", "/api/projects/demo/logs/stream") as resp:
            lines = []
            for line in resp.iter_lines():
                lines.append(line)
                if "[SYSTEM]" in line:
                    break
        assert any("[SYSTEM]" in line for line in lines)

    def test_buffered_logs_are_streamed(self, client, tmp_path, monkeypatch):
        """파이프라인 실행 중 쌓인 로그가 스트림에 포함된다."""
        import app.main as main_module
        # Deprecated PROJECT_ROOT patch removed
        _pipeline_jobs.clear()
        _job_set("demo", {"status": "done"})
        main_module._project_logs["demo"] = ["line1", "line2"]

        with client.stream("GET", "/api/projects/demo/logs/stream") as resp:
            content = b"".join(resp.iter_bytes()).decode()

        assert "line1" in content
        assert "line2" in content
        main_module._project_logs.pop("demo", None)


# ─────────────────────────────────────────────────────────────────────────────
# DELETE /api/projects/{name} — 프로젝트 삭제
# ─────────────────────────────────────────────────────────────────────────────

class TestDeleteProject:
    def test_invalid_name_returns_400(self, client):
        resp = client.delete("/api/projects/../delete")
        assert resp.status_code in (400, 404)

    def test_delete_nonexistent_project_succeeds_gracefully(self, client):
        resp = client.delete("/api/projects/nonexistent_test")
        assert resp.status_code == 200
        assert resp.json()["status"] == "success"

    def test_complete_deletion_removes_files_and_manifest_entry(self, client, tmp_path):
        import app.main as main_module
        
        # 임시 환경 모킹 설정
        proj_name = "test_del"
        proj_dir = tmp_path / "storage" / "projects" / proj_name
        proj_dir.mkdir(parents=True)
        (proj_dir / "config.yaml").write_text("project: test_del", encoding="utf-8")
        
        # 2. 배포용 JSON 생성
        data_dir = tmp_path / "frontend" / "public" / "data"
        data_dir.mkdir(parents=True)
        frontend_json = data_dir / f"{proj_name}_data.json"
        frontend_json.write_text("{}", encoding="utf-8")
        
        # 3. 매니페스트 생성 및 엔트리 삽입
        manifest_path = data_dir / "projects.json"
        manifest_data = [
            {"id": "other_proj", "name": "다른프로젝트"},
            {"id": proj_name, "name": "삭제대상프로젝트"}
        ]
        manifest_path.write_text(json.dumps(manifest_data), encoding="utf-8")
        
        # API 호출
        resp = client.delete(f"/api/projects/{proj_name}")
        assert resp.status_code == 200
        assert resp.json()["status"] == "success"
        
        # 검증: 폴더 삭제되었는가?
        assert not proj_dir.exists()
        
        # 검증: 배포 JSON 삭제되었는가?
        assert not frontend_json.exists()
        
        # 검증: 매니페스트에서 제거되었는가?
        with open(manifest_path, encoding="utf-8") as f:
            updated_manifest = json.load(f)
        assert len(updated_manifest) == 1
        assert updated_manifest[0]["id"] == "other_proj"


class TestUpdateProjectsManifest:
    """_update_projects_manifest()가 기존 published 값을 유지하는지 검증.

    회귀 버그: 기존 항목을 리스트에서 먼저 제거한 뒤 같은 리스트에서 재검색해
    existing이 항상 {}가 되어, 호출할 때마다 published가 무조건 False로
    리셋되던 문제(main.py:163-172).
    """

    def test_preserves_published_true_on_repeated_calls(self, tmp_path, monkeypatch):
        import app.main as main_module
        monkeypatch.setattr(main_module, "FRONTEND_ROOT", tmp_path / "frontend")

        manifest_path = tmp_path / "frontend" / "public" / "data" / "projects.json"
        manifest_path.parent.mkdir(parents=True)
        manifest_path.write_text(
            json.dumps([{"id": "mumhwa", "name": "mumhwa", "file": "mumhwa_data.json",
                         "updated": "2026-07-01 00:00", "published": True}]),
            encoding="utf-8",
        )

        # 파이프라인을 다시 실행/내보내기할 때마다 호출되는 상황을 재현
        main_module._update_projects_manifest("mumhwa", "mumhwa_data.json")
        main_module._update_projects_manifest("mumhwa", "mumhwa_data.json")

        with open(manifest_path, encoding="utf-8") as f:
            projects = json.load(f)
        entry = next(p for p in projects if p["id"] == "mumhwa")
        assert entry["published"] is True

    def test_new_project_defaults_to_unpublished(self, tmp_path, monkeypatch):
        import app.main as main_module
        monkeypatch.setattr(main_module, "FRONTEND_ROOT", tmp_path / "frontend")

        main_module._update_projects_manifest("new_proj", "new_proj_data.json")

        manifest_path = tmp_path / "frontend" / "public" / "data" / "projects.json"
        with open(manifest_path, encoding="utf-8") as f:
            projects = json.load(f)
        entry = next(p for p in projects if p["id"] == "new_proj")
        assert entry["published"] is False


# ─────────────────────────────────────────────────────────────────────────────
# POST /api/projects/{name}/preview — 정제 규칙 미리보기
# ─────────────────────────────────────────────────────────────────────────────

class TestPreviewProjectConfig:
    """미리보기가 실제 파이프라인(run)과 동일하게 addr_split 같은 다중 출력
    transform을 처리하는지 검증한다. 회귀 버그: address_parsing이 주입되지
    않아 addr_split이 항상 빈 값을 반환하고, 파생열(_시도 등)이 미리보기에
    아예 나타나지 않던 문제(main.py의 preview_project_config)."""

    def test_addr_split_produces_derived_columns_with_values(self, client, tmp_path):
        proj_dir = tmp_path / "storage" / "projects" / "demo"
        proj_dir.mkdir(parents=True)

        data_xlsx = tmp_path / "주소원본.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.cell(1, 1, "주소")
        ws.cell(2, 1, "서울특별시 강남구 테헤란로 123")
        wb.save(data_xlsx)

        cfg = {
            "project": "demo",
            "source": {
                "file": str(data_xlsx),
                "sheet": None,
                "header_row": 1,
                "data_start_row": 2,
            },
            "paths": {"output_dir": "output", "output_file": "result.xlsx"},
            "sheets": {"cleaned": "Cleaned", "summary": "Summary"},
            "address_parsing": {
                "sido_patterns": [["서울", ["서울", "서울특별시"]]],
                "seoul_gu": ["강남구"],
            },
            "columns": [
                {"output_col": "주소", "source_col": 1, "transform": "addr_split"},
            ],
        }
        (proj_dir / "config.yaml").write_text(yaml.dump(cfg, allow_unicode=True), encoding="utf-8")

        resp = client.post("/api/projects/demo/preview")

        assert resp.status_code == 200
        body = resp.json()
        row = body["preview"][0]
        assert row["cleaned"]["주소_시도"] == "서울"
        assert row["cleaned"]["주소_시군구"] == "강남구"
        # 파생열이 raw 쪽에도 같은 키로 존재해야 프론트 렌더링(raw/cleaned 동일 키 순회)과 맞는다
        assert "주소_시도" in row["raw"]


# ─────────────────────────────────────────────────────────────────────────────
# _build_single_html_template — 독립 실행형 HTML 리포트
# ─────────────────────────────────────────────────────────────────────────────

class TestBuildSingleHtmlTemplate:
    """단독 HTML 리포트가 실제로 오프라인에서 동작하는지(CDN 미의존) 및
    XSS(스크립트 컨텍스트 탈출)로부터 안전한지 검증한다."""

    def _sample_data(self) -> dict:
        return {
            "project": "demo",
            "rows": [{"이름": "홍길동", "부서": "개발팀"}],
            "columns": [{"key": "이름"}, {"key": "부서"}],
            "dashboard": {"kpi": [], "charts": [], "list": {"visible_cols": ["이름", "부서"], "filter_cols": []}},
        }

    def test_no_external_cdn_references(self):
        html_doc = _build_single_html_template("demo", self._sample_data())
        for banned in ("cdn.tailwindcss.com", "cdn.jsdelivr.net", "fonts.googleapis.com"):
            assert banned not in html_doc

    def test_embeds_vendored_libraries(self):
        html_doc = _build_single_html_template("demo", self._sample_data())
        assert "Chart.js v4.5.1" in html_doc
        assert "SheetJS" in html_doc
        assert "tailwindcss" in html_doc.lower()

    def test_project_name_is_html_escaped(self):
        html_doc = _build_single_html_template("<img src=x onerror=alert(1)>", self._sample_data())
        assert "<img src=x onerror=alert(1)>" not in html_doc
        assert "&lt;img src=x onerror=alert(1)&gt;" in html_doc

    def test_data_containing_script_close_tag_cannot_break_out(self):
        data = self._sample_data()
        data["rows"] = [{"이름": "</script><script>alert(1)</script>", "부서": "개발팀"}]
        html_doc = _build_single_html_template("demo", data)
        assert "</script><script>alert(1)</script>" not in html_doc
        assert "u003c/script" in html_doc

    def test_produces_well_formed_document(self):
        html_doc = _build_single_html_template("demo", self._sample_data())
        assert html_doc.startswith("<!DOCTYPE html>")
        assert html_doc.rstrip().endswith("</html>")


# ─────────────────────────────────────────────────────────────────────────────
# POST /api/projects/create-merge
# ─────────────────────────────────────────────────────────────────────────────

class TestCreateMergeProject:
    """복수 엑셀 파일 업로드 → 병합 → draft 프로젝트 생성 엔드포인트 검증."""

    def _make_source_xlsx(self, path: Path, rows: list[tuple[str, str]]) -> None:
        wb = Workbook()
        ws = wb.active
        ws.cell(1, 1, "이름")
        ws.cell(1, 2, "부서")
        for i, (name, dept) in enumerate(rows, start=2):
            ws.cell(i, 1, name)
            ws.cell(i, 2, dept)
        wb.save(path)

    def _post_merge(
        self,
        client,
        tmp_path: Path,
        project_name: str = "merged_demo",
        options: dict | None = None,
        num_files: int = 2,
    ):
        src_paths = []
        for idx in range(num_files):
            p = tmp_path / f"source_{idx}.xlsx"
            self._make_source_xlsx(p, [(f"홍길동{idx}", f"팀{idx}")])
            src_paths.append(p)

        opts = options if options is not None else {
            "dedup_strategy": "none",
            "key_cols": [],
            "add_source_col": True,
            "source_col_name": "_출처파일",
        }

        files = [
            ("files", (p.name, open(p, "rb"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
            for p in src_paths
        ]
        data = {"name": project_name, "options": json.dumps(options if options is not None else opts)}
        try:
            return client.post("/api/projects/create-merge", data=data, files=files)
        finally:
            for _, (_, fh, _) in files:
                fh.close()

    def test_merge_two_files_returns_success(self, client, tmp_path):
        resp = self._post_merge(client, tmp_path)
        assert resp.status_code == 200
        body = resp.json()
        assert body["status"] == "success"
        assert body["project"] == "merged_demo"
        assert "draft_path" in body
        assert "merged_path" in body

    def test_merge_creates_project_draft_and_config(self, client, tmp_path):
        import app.main as main_module
        resp = self._post_merge(client, tmp_path, project_name="merged_check")
        assert resp.status_code == 200

        proj_dir = main_module.STORAGE_ROOT / "projects" / "merged_check"
        assert (proj_dir / "config.yaml").exists()
        config_data = yaml.safe_load((proj_dir / "config.yaml").read_text(encoding="utf-8"))
        assert "merge" in config_data
        assert len(config_data["merge"]["sources"]) == 2

    def test_single_file_rejected(self, client, tmp_path):
        resp = self._post_merge(client, tmp_path, num_files=1)
        assert resp.status_code == 400

    def test_invalid_project_name_rejected(self, client, tmp_path):
        resp = self._post_merge(client, tmp_path, project_name="../etc")
        assert resp.status_code == 400

    def test_malformed_options_json_rejected(self, client, tmp_path):
        src_paths = []
        for idx in range(2):
            p = tmp_path / f"source_{idx}.xlsx"
            self._make_source_xlsx(p, [(f"홍길동{idx}", f"팀{idx}")])
            src_paths.append(p)
        files = [
            ("files", (p.name, open(p, "rb"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
            for p in src_paths
        ]
        try:
            resp = client.post(
                "/api/projects/create-merge",
                data={"name": "bad_opts", "options": "{not valid json"},
                files=files,
            )
        finally:
            for _, (_, fh, _) in files:
                fh.close()
        assert resp.status_code == 400
