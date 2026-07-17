"""
Tests for 오버레이 ↔ 파이프라인 연결 (dashboard_edit_plan §4.1 · 4단계)

**이 계획이 존재하는 이유가 여기서 검증된다.**

`output/<name>_cleaned.xlsx` 는 `storage/raw` + `config.yaml` 에서 매번 새로 만드는 파생물이다.
편집을 거기 직접 쓰면 다음 `/run` 이 지운다. 오버레이는 그걸 막기 위한 것이므로,
**정제를 몇 번을 다시 돌려도 편집이 살아남아야 한다.**

실행: pytest tests/test_overrides_pipeline.py -v
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from engine.config import ROW_ID_COL, SurveyConfig
from engine.exporter import export_to_json
from engine.overrides import (
    REASON_COL_MISSING,
    REASON_PREV_MISMATCH,
    REASON_ROW_MISSING,
    Edit,
    Overrides,
    save_overrides,
)
from engine.pipeline import SurveyPipeline


@pytest.fixture
def proj(tmp_path: Path) -> Path:
    """storage/projects/<name>/ 를 흉내낸 프로젝트 폴더 + 원본 엑셀."""
    d = tmp_path / "proj"
    d.mkdir()
    p = d / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(1, 1, "이름"); ws.cell(1, 2, "지역")
    rows = {2: ("홍길동", "서울"), 3: ("김철수", "부산"), 5: ("이영희", "서울")}  # 4행은 빈 행
    for r, (n, g) in rows.items():
        ws.cell(r, 1, n); ws.cell(r, 2, g)
    wb.save(p)
    return d


def _cfg(proj: Path) -> SurveyConfig:
    cfg = SurveyConfig.model_validate(dict(
        project="ov",
        source={"file": "data.xlsx", "sheet": "Sheet1", "header_row": 1},
        sheets={"cleaned": "Cleaned", "summary": "Summary"},
        columns=[
            {"output_col": "이름", "source_col": 1, "transform": "copy"},
            {"output_col": "지역", "source_col": 2, "transform": "copy"},
        ],
        paths={"output_dir": str(proj / "out"), "output_file": "r.xlsx"},
        excel_options={"include_slicers": False, "include_charts": False},
    ))
    return cfg


def _run(proj: Path) -> tuple[Path, SurveyConfig, SurveyPipeline]:
    cfg = _cfg(proj)
    # _config_path 가 proj_dir 을 정한다 — 파이프라인은 여기서 overrides.json 을 찾는다
    pipe = SurveyPipeline(cfg, config_path=proj / "config.yaml")
    out = pipe.run(input_path=proj / "data.xlsx")
    return out, cfg, pipe


def _cleaned_col(out: Path, cfg: SurveyConfig, col: str) -> list:
    ws = openpyxl.load_workbook(out)[cfg.sheets.cleaned]
    rows = list(ws.rows)
    headers = [c.value for c in rows[1]]
    i = headers.index(col)
    return [r[i].value for r in rows[2:]]


# ─────────────────────────────────────────────────────────────────────────────
# 핵심: 편집이 재실행에서 살아남는가
# ─────────────────────────────────────────────────────────────────────────────

class TestEditSurvivesRerun:
    def test_no_overrides_is_unchanged(self, proj):
        out, cfg, _ = _run(proj)
        assert _cleaned_col(out, cfg, "지역") == ["서울", "부산", "서울"]

    def test_hand_edit_overrides_transform_output(self, proj):
        save_overrides(proj, Overrides(edits=[
            Edit(row_id="r3", col="지역", value="부산광역시", prev="부산", has_prev=True),
        ]))
        out, cfg, _ = _run(proj)
        assert _cleaned_col(out, cfg, "지역") == ["서울", "부산광역시", "서울"]

    def test_edit_survives_repeated_runs(self, proj):
        """★ 이 계획의 존재 이유 — /run 을 몇 번 돌려도 편집이 살아 있어야 한다."""
        save_overrides(proj, Overrides(edits=[
            Edit(row_id="r3", col="지역", value="부산광역시"),
        ]))
        for _ in range(3):
            out, cfg, _ = _run(proj)
            assert _cleaned_col(out, cfg, "지역") == ["서울", "부산광역시", "서울"]

    def test_edit_lands_on_right_row_despite_blank_rows(self, proj):
        """빈 행(4행)이 있어도 r5 는 원본 5행 = 이영희에 붙어야 한다.

        순번 기반이었다면 r5 는 존재하지 않거나 엉뚱한 행을 가리켰을 것이다.
        """
        save_overrides(proj, Overrides(edits=[
            Edit(row_id="r5", col="지역", value="서울특별시"),
        ]))
        out, cfg, _ = _run(proj)
        assert _cleaned_col(out, cfg, "지역") == ["서울", "부산", "서울특별시"]
        assert _cleaned_col(out, cfg, "이름") == ["홍길동", "김철수", "이영희"]

    def test_multiple_cols_same_row(self, proj):
        save_overrides(proj, Overrides(edits=[
            Edit(row_id="r2", col="이름", value="홍길순"),
            Edit(row_id="r2", col="지역", value="서울특별시"),
        ]))
        out, cfg, _ = _run(proj)
        assert _cleaned_col(out, cfg, "이름")[0] == "홍길순"
        assert _cleaned_col(out, cfg, "지역")[0] == "서울특별시"

    def test_edit_reaches_dashboard_json_and_aggregates(self, proj):
        """편집이 xlsx 를 지나 대시보드 데이터·집계까지 도달해야 한다."""
        save_overrides(proj, Overrides(edits=[
            Edit(row_id="r3", col="지역", value="서울"),      # 부산 → 서울
        ]))
        out, cfg, _ = _run(proj)
        data = export_to_json(out, cfg)
        assert [r["지역"] for r in data["rows"]] == ["서울", "서울", "서울"]
        assert data["aggregates"]["지역"] == {"서울": 3}      # 차트가 이 값을 그린다

    def test_override_is_reflected_in_summary_source_values(self, proj):
        """엑셀 Summary 시트 집계도 편집을 반영해야 한다 (cleaned_col_vals 경유)."""
        from engine.pipeline import _build_registry
        from engine.writer import CleanedSheetWriter
        from engine.overrides import OverrideApplier
        import openpyxl as _o

        cfg = _cfg(proj)
        wb = _o.load_workbook(proj / "data.xlsx", data_only=True)
        from engine.pipeline import _sheet_to_dataframe
        df = _sheet_to_dataframe(wb, cfg)
        ap = OverrideApplier(Overrides(edits=[Edit(row_id="r3", col="지역", value="서울")]))
        out_wb = Workbook(); out_wb.remove(out_wb.active)
        _, _, cleaned_col_vals = CleanedSheetWriter(cfg, _build_registry(cfg)).write(out_wb, df, ap)
        assert cleaned_col_vals["지역"] == ["서울", "서울", "서울"]


# ─────────────────────────────────────────────────────────────────────────────
# 충돌 보고 — 편집이 조용히 사라지지 않는가
# ─────────────────────────────────────────────────────────────────────────────

class TestConflictReporting:
    def test_clean_run_has_no_conflicts(self, proj):
        save_overrides(proj, Overrides(edits=[Edit(row_id="r3", col="지역", value="부산광역시")]))
        _, _, pipe = _run(proj)
        assert pipe.override_conflicts == []

    def test_missing_row_is_reported_not_dropped(self, proj):
        """원본에 없는 행을 가리키는 편집 — 버리지 않고 보고한다."""
        save_overrides(proj, Overrides(edits=[Edit(row_id="r99", col="지역", value="제주")]))
        _, _, pipe = _run(proj)
        (c,) = pipe.override_conflicts
        assert c.reason == REASON_ROW_MISSING
        assert c.row_id == "r99"
        assert c.value == "제주"

    def test_missing_col_is_reported(self, proj):
        save_overrides(proj, Overrides(edits=[Edit(row_id="r2", col="없는컬럼", value="x")]))
        _, _, pipe = _run(proj)
        (c,) = pipe.override_conflicts
        assert c.reason == REASON_COL_MISSING

    def test_prev_mismatch_reported_but_applied(self, proj):
        """원본이 바뀌어 산출값이 달라진 경우 — 편집은 적용하되 사실을 알린다."""
        save_overrides(proj, Overrides(edits=[
            Edit(row_id="r3", col="지역", value="부산광역시", prev="대전", has_prev=True),
        ]))
        out, cfg, pipe = _run(proj)
        assert _cleaned_col(out, cfg, "지역")[1] == "부산광역시"   # 적용됨
        (c,) = pipe.override_conflicts
        assert c.reason == REASON_PREV_MISMATCH
        assert c.expected_prev == "대전"
        assert c.actual_prev == "부산"


# ─────────────────────────────────────────────────────────────────────────────
# 행 필터와의 상호작용 — 2단계 회귀 방지의 실전판
# ─────────────────────────────────────────────────────────────────────────────

class TestWithRowFilter:
    def test_edit_lands_correctly_when_rows_filtered_out(self, proj):
        """행 필터로 앞 행이 빠져도 편집은 제 행에 붙어야 한다.

        preprocessor 가 index 를 리셋하면 r5 가 다른 행에 붙는다 — 그 회귀의 실전 검증.
        """
        from engine.config import RowFilterRule

        cfg = _cfg(proj)
        # r3(부산) 제외 → r2, r5 만 남는다
        cfg.preprocess.row_filter.exclude = [RowFilterRule(col=2, equals="부산")]
        save_overrides(proj, Overrides(edits=[Edit(row_id="r5", col="이름", value="이영희님")]))
        pipe = SurveyPipeline(cfg, config_path=proj / "config.yaml")
        out = pipe.run(input_path=proj / "data.xlsx")

        assert _cleaned_col(out, cfg, "이름") == ["홍길동", "이영희님"]
        assert pipe.override_conflicts == []
        ws = openpyxl.load_workbook(out)[cfg.sheets.cleaned]
        rids = [r[-1].value for r in list(ws.rows)[2:]]
        assert rids == ["r2", "r5"]
