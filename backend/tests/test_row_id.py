"""
Tests for `__row_id` 파이프라인 관통 (dashboard_edit_plan §2 · 2단계)

`__row_id` 는 대시보드 편집이 "몇 번째 행"이 아니라 **"어느 행"**에 붙는지 고정하는 숨은 열이다.
원본 엑셀의 실제 행번호(빈 행을 건너뛰기 "전" 기준)에서 나온다.

**이게 틀어지면 편집이 조용히 엉뚱한 행에 적용된다** — 그래서 다음을 못박는다:
  1. 빈 행을 건너뛰어도 행번호는 원본 기준을 유지하는가
  2. 행 필터(row_filter)를 켜도 살아남은 행의 id 가 보존되는가  ← reset_index 회귀 방지
  3. Cleaned 시트에 숨김 열로 기록되고 Excel Table 안에 들어가는가
  4. 사용자 눈(meta.columns)에는 안 보이고 rows 에는 있는가
  5. source_col 이 __row_id 를 데이터로 읽어가지 않는가

실행: pytest tests/test_row_id.py -v
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from engine.config import ROW_ID_COL, SurveyConfig
from engine.exporter import build_data_json, export_to_json
from engine.pipeline import SurveyPipeline, _sheet_to_dataframe
from engine.preprocessor import Preprocessor


def _cfg(tmp_path: Path, **over) -> SurveyConfig:
    base = dict(
        project="rid_proj",
        source={"file": "data.xlsx", "sheet": "Sheet1", "header_row": 1},
        sheets={"cleaned": "Cleaned", "summary": "Summary"},
        columns=[
            {"output_col": "이름", "source_col": 1, "transform": "copy"},
            {"output_col": "부서", "source_col": 2, "transform": "copy"},
        ],
        paths={"output_dir": str(tmp_path / "out"), "output_file": "r.xlsx"},
        excel_options={"include_slicers": False, "include_charts": False},
    )
    base.update(over)
    return SurveyConfig.model_validate(base)


def _raw_xlsx_with_gaps(tmp_path: Path) -> Path:
    """2행부터 데이터. 4행과 7행을 일부러 비운다 → 행번호에 구멍이 생긴다."""
    p = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(1, 1, "이름"); ws.cell(1, 2, "부서")
    layout = {
        2: ("홍길동", "개발팀"),
        3: ("김철수", "영업팀"),
        # 4행: 빈 행
        5: ("이영희", "개발팀"),
        6: ("박민준", "영업팀"),
        # 7행: 빈 행
        8: ("최지훈", "기획팀"),
    }
    for r, (name, dept) in layout.items():
        ws.cell(r, 1, name); ws.cell(r, 2, dept)
    wb.save(p)
    return p


# ─────────────────────────────────────────────────────────────────────────────
# 1·2. 행번호가 원본 기준으로 살아남는가
# ─────────────────────────────────────────────────────────────────────────────

class TestRowIdSurvives:
    def test_blank_rows_do_not_shift_row_ids(self, tmp_path):
        """빈 행을 건너뛰어도 index 는 원본 엑셀 행번호여야 한다."""
        wb = openpyxl.load_workbook(_raw_xlsx_with_gaps(tmp_path))
        df = _sheet_to_dataframe(wb, _cfg(tmp_path))
        # 4·7행이 비었으므로 2,3,5,6,8 이어야 한다 (0,1,2,3,4 가 아니라)
        assert list(df.index) == [2, 3, 5, 6, 8]
        assert df.index.name == ROW_ID_COL

    def test_row_filter_preserves_row_ids(self, tmp_path):
        """행 필터가 index 를 리셋하면 안 된다.

        preprocessor 의 reset_index(drop=True) 회귀 방지 — 이게 돌아오면 행 필터를 쓰는
        프로젝트에서 편집이 통째로 엉뚱한 행에 붙는다.
        """
        wb = openpyxl.load_workbook(_raw_xlsx_with_gaps(tmp_path))
        df = _sheet_to_dataframe(wb, _cfg(tmp_path))
        # 부서(2번 열)가 '영업팀' 인 행 제외 → 3,6 이 빠지고 2,5,8 만 남아야 한다
        cfg = _cfg(tmp_path, preprocess={
            "row_filter": {"exclude": [{"col": 2, "equals": "영업팀"}]}
        })
        out = Preprocessor(cfg.preprocess).run(df)
        assert list(out.index) == [2, 5, 8]

    def test_empty_sheet_returns_empty_frame(self, tmp_path):
        p = tmp_path / "empty.xlsx"
        wb = Workbook(); wb.active.title = "Sheet1"; wb.active.cell(1, 1, "이름"); wb.save(p)
        df = _sheet_to_dataframe(openpyxl.load_workbook(p), _cfg(tmp_path))
        assert df.empty


# ─────────────────────────────────────────────────────────────────────────────
# 3. Cleaned 시트에 어떻게 들어가는가
# ─────────────────────────────────────────────────────────────────────────────

class TestRowIdInCleanedSheet:
    @pytest.fixture
    def built(self, tmp_path):
        _raw_xlsx_with_gaps(tmp_path)
        cfg = _cfg(tmp_path)
        out = SurveyPipeline(cfg).run(input_path=tmp_path / "data.xlsx")
        return out, cfg

    def test_row_id_is_last_column_and_hidden(self, built):
        out, cfg = built
        ws = openpyxl.load_workbook(out)[cfg.sheets.cleaned]
        headers = [c.value for c in list(ws.rows)[1]]
        assert headers == ["이름", "부서", ROW_ID_COL]

        from openpyxl.utils import get_column_letter
        letter = get_column_letter(len(headers))
        assert ws.column_dimensions[letter].hidden is True

    def test_row_id_values_match_original_excel_rows(self, built):
        out, cfg = built
        ws = openpyxl.load_workbook(out)[cfg.sheets.cleaned]
        rows = list(ws.rows)[2:]                       # 데이터는 3행부터
        rids = [r[-1].value for r in rows]
        assert rids == ["r2", "r3", "r5", "r6", "r8"]  # 빈 행(4·7) 건너뛴 원본 번호

    def test_row_id_is_inside_excel_table(self, built):
        """사용자가 엑셀에서 정렬·필터해도 식별자가 제 행에 붙어 있어야 한다."""
        out, cfg = built
        ws = openpyxl.load_workbook(out)[cfg.sheets.cleaned]
        ref = list(ws.tables.values())[0].ref          # 예: "A2:C7"
        last_col_letter = ref.split(":")[1].rstrip("0123456789")
        from openpyxl.utils import get_column_letter
        assert last_col_letter == get_column_letter(3)  # 이름·부서·__row_id


# ─────────────────────────────────────────────────────────────────────────────
# 4. 사용자 눈에 보이지 않는가
# ─────────────────────────────────────────────────────────────────────────────

class TestRowIdHiddenFromDashboard:
    @pytest.fixture
    def exported(self, tmp_path):
        _raw_xlsx_with_gaps(tmp_path)
        cfg = _cfg(tmp_path)
        out = SurveyPipeline(cfg).run(input_path=tmp_path / "data.xlsx")
        return export_to_json(out, cfg)

    def test_row_id_present_in_rows(self, exported):
        """편집이 행을 찾으려면 rows 에는 반드시 있어야 한다."""
        assert [r[ROW_ID_COL] for r in exported["rows"]] == ["r2", "r3", "r5", "r6", "r8"]

    def test_row_id_absent_from_meta_columns(self, exported):
        """표·차트·필터·드로어·CSV 는 meta.columns 로 그려진다 — 여기 없으면 안 보인다."""
        assert ROW_ID_COL not in {c["key"] for c in exported["meta"]["columns"]}

    def test_row_id_absent_from_aggregates(self, exported):
        assert ROW_ID_COL not in exported["aggregates"]

    def test_row_id_absent_from_default_dashboard(self, exported):
        """자동 생성 대시보드가 __row_id 로 차트를 만들면 안 된다."""
        dash = exported["dashboard"]
        refs = {c.get("col") for c in dash["charts"]}
        refs |= set(dash["list"]["visible_cols"]) | set(dash["list"]["filter_cols"])
        assert ROW_ID_COL not in refs

    def test_build_data_json_excludes_row_id_from_columns(self):
        """편집 경로(headers 미지정)에서도 동일해야 한다."""
        rows = [
            {"부서": "개발팀", ROW_ID_COL: "r2"},
            {"부서": "영업팀", ROW_ID_COL: "r3"},
            {"부서": "개발팀", ROW_ID_COL: "r5"},
            {"부서": "영업팀", ROW_ID_COL: "r6"},
        ]
        cfg = SurveyConfig(
            project="p", sheets={"cleaned": "Cleaned", "summary": "Summary"},
            columns=[], paths={"output_dir": "o", "output_file": "f.xlsx"},
        )
        res = build_data_json(rows, cfg)
        assert ROW_ID_COL not in {c["key"] for c in res["meta"]["columns"]}
        assert ROW_ID_COL not in res["aggregates"]
        assert res["rows"][0][ROW_ID_COL] == "r2"      # rows 에는 남아야 한다


# ─────────────────────────────────────────────────────────────────────────────
# 5. 안전장치
# ─────────────────────────────────────────────────────────────────────────────

class TestRowIdSafety:
    def test_reserved_name_collision_raises(self, tmp_path):
        """사용자 컬럼명이 __row_id 면 조용히 덮이는 대신 즉시 실패해야 한다."""
        _raw_xlsx_with_gaps(tmp_path)
        cfg = _cfg(tmp_path, columns=[
            {"output_col": "이름", "source_col": 1, "transform": "copy"},
            {"output_col": ROW_ID_COL, "source_col": 2, "transform": "copy"},
        ])
        with pytest.raises(ValueError, match=ROW_ID_COL):
            SurveyPipeline(cfg).run(input_path=tmp_path / "data.xlsx")

    def test_source_col_out_of_range_does_not_read_row_id(self, tmp_path):
        """행번호를 index 로 둔 이유 — 컬럼이었다면 범위 밖 source_col 이 이걸 읽어간다."""
        _raw_xlsx_with_gaps(tmp_path)
        cfg = _cfg(tmp_path, columns=[
            {"output_col": "이름", "source_col": 1, "transform": "copy"},
            {"output_col": "없는열", "source_col": 99, "transform": "copy"},
        ])
        out = SurveyPipeline(cfg).run(input_path=tmp_path / "data.xlsx")
        ws = openpyxl.load_workbook(out)[cfg.sheets.cleaned]
        vals = [r[1].value for r in list(ws.rows)[2:]]   # '없는열'
        assert all(v is None for v in vals)              # 행번호가 새어들면 안 된다
