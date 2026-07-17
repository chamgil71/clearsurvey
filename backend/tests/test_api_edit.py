"""
Tests for 편집 API (dashboard_edit_plan §4.4 · §5.1 · 5단계)

  GET    /api/projects/{name}/overrides
  PATCH  /api/projects/{name}/rows/{row_id}
  DELETE /api/projects/{name}/overrides
  POST   /api/projects/{name}/rebuild
  GET    /api/projects/{name}/freshness   (확장)
  GET    /api/projects/{name}/download    (뒤처지면 rebuild)

못박는 것 — 계획 §5.1 의 순서가 실제로 그 순서인가:
  ④ overrides 갱신 → ⑤ rows 패치 → ⑥ 재계산 → ⑦ 저장·복사 → ⑧ 응답
  그리고 **저장이 xlsx 를 건드리지 않는가**(지연 생성의 존재 이유)

실행: pytest tests/test_api_edit.py -v
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml
from fastapi.testclient import TestClient
from openpyxl import Workbook

_PROJECT_ROOT = Path(__file__).parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from engine.config import ROW_ID_COL


NAME = "edit_proj"


@pytest.fixture
def client(tmp_path, monkeypatch):
    import app.main as main_module

    monkeypatch.setattr(main_module, "STORAGE_ROOT", tmp_path / "storage")
    monkeypatch.setattr(main_module, "FRONTEND_ROOT", tmp_path / "frontend")
    (tmp_path / "frontend" / "public" / "data").mkdir(parents=True)
    return TestClient(main_module.app)


@pytest.fixture
def proj(tmp_path):
    """정제·내보내기까지 끝난 프로젝트를 만들어 둔다 (편집의 출발선)."""
    d = tmp_path / "storage" / "projects" / NAME
    d.mkdir(parents=True)

    src = d / "data.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    ws.cell(1, 1, "이름"); ws.cell(1, 2, "지역"); ws.cell(1, 3, "기관코드")
    # 값이 반복돼야 지역이 category 로 잡힌다.
    # 기관코드는 **텍스트 서식의 숫자 문자열** — 타입 고정(§5.1 ⑥)이 필요한 바로 그 컬럼.
    rows = {
        2: ("홍길동", "서울", "3000000"), 3: ("김철수", "부산", "6520000"),
        5: ("이영희", "서울", "3000000"), 6: ("박민준", "부산", "6520000"),
        7: ("최지훈", "서울", "3000000"), 9: ("정수빈", "대구", "6520000"),
    }
    for r, (n, g, c) in rows.items():
        ws.cell(r, 1, n); ws.cell(r, 2, g); ws.cell(r, 3, c)
    wb.save(src)

    cfg = {
        "project": NAME,
        "source": {"file": "data.xlsx", "sheet": "Sheet1", "header_row": 1},
        "sheets": {"cleaned": "Cleaned", "summary": "Summary"},
        "columns": [
            {"output_col": "이름", "source_col": 1, "transform": "copy"},
            {"output_col": "지역", "source_col": 2, "transform": "copy"},
            {"output_col": "기관코드", "source_col": 3, "transform": "copy"},
        ],
        "paths": {"output_dir": "out", "output_file": "r.xlsx"},
        "excel_options": {"include_slicers": False, "include_charts": False},
    }
    (d / "config.yaml").write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    return d


def _seed(client):
    """run + export 를 끝내 data.json 을 만든다."""
    r = client.post(f"/api/projects/{NAME}/rebuild")
    assert r.status_code == 200, r.text
    return r


def _data(client):
    return client.get(f"/api/projects/{NAME}/overrides")


def _xlsx_mtime(proj) -> float:
    return (proj / "out" / "r.xlsx").stat().st_mtime


# ─────────────────────────────────────────────────────────────────────────────
# 저장 (PATCH /rows)
# ─────────────────────────────────────────────────────────────────────────────

class TestPatchRow:
    def test_saves_edit_and_returns_updated_data(self, client, proj):
        _seed(client)
        r = client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["status"] == "success"
        assert body["edit_count"] == 1

        # ⑧ 갱신된 데이터가 응답에 실려 온다 — 프런트가 이걸로 교체한다
        rows = body["data"]["rows"]
        assert next(x["지역"] for x in rows if x[ROW_ID_COL] == "r3") == "부산광역시"

    def test_aggregates_recomputed_so_charts_update(self, client, proj):
        _seed(client)
        r = client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "서울"})
        agg = r.json()["data"]["aggregates"]["지역"]
        assert agg == {"서울": 4, "부산": 1, "대구": 1}      # 부산 2 → 1, 서울 3 → 4

    def test_does_not_touch_xlsx(self, client, proj):
        """★ 지연 생성의 존재 이유 — 저장은 xlsx 를 건드리지 않는다."""
        _seed(client)
        before = _xlsx_mtime(proj)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        assert _xlsx_mtime(proj) == before

    def test_writes_both_local_and_published_json(self, client, proj, tmp_path):
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        published = tmp_path / "frontend" / "public" / "data" / f"{NAME}_data.json"
        local = proj / f"{NAME}_data.json"
        assert published.exists() and local.exists()
        pub = json.loads(published.read_text(encoding="utf-8"))
        assert next(x["지역"] for x in pub["rows"] if x[ROW_ID_COL] == "r3") == "부산광역시"

    def test_multiple_cols_one_request(self, client, proj):
        _seed(client)
        r = client.patch(f"/api/projects/{NAME}/rows/r2", json={"이름": "홍길순", "지역": "인천"})
        assert r.json()["edit_count"] == 2          # 행 1건 저장 = 컬럼별 편집 2건

    def test_edit_does_not_flip_numeric_string_column_type(self, client, proj):
        """★ 저장이 컬럼의 타입을 바꾸면 안 된다 (§5.1 ⑥).

        `기관코드` 는 텍스트 서식의 숫자 문자열이라 xlsx 원시값으로는 category 지만,
        data.json 의 rows 에는 clean_value 를 거쳐 정수로 실린다. PATCH 가 타입을
        물려주지 않고 재감지하면 numeric 으로 뒤집혀 **필터 드롭다운이 사라진다.**
        """
        _seed(client)
        before = client.get(f"/api/projects/{NAME}/freshness")  # noqa: F841
        base = json.loads((proj / f"{NAME}_data.json").read_text(encoding="utf-8"))
        assert next(c["type"] for c in base["meta"]["columns"] if c["key"] == "기관코드") == "category"
        assert "기관코드" in base["aggregates"]

        r = client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        data = r.json()["data"]

        kind = next(c["type"] for c in data["meta"]["columns"] if c["key"] == "기관코드")
        assert kind == "category", "저장이 기관코드를 numeric 으로 뒤집었다"
        assert "기관코드" in data["aggregates"], "기관코드 필터 드롭다운이 사라졌다"

    def test_repeated_saves_do_not_drift_types(self, client, proj):
        """저장은 여러 번 일어난다 — 타입이 서서히 흘러내리면 안 된다."""
        _seed(client)
        for i in range(4):
            r = client.patch(f"/api/projects/{NAME}/rows/r2", json={"이름": f"홍길동{i}"})
        data = r.json()["data"]
        assert next(c["type"] for c in data["meta"]["columns"] if c["key"] == "기관코드") == "category"
        assert "기관코드" in data["aggregates"]

    def test_new_value_appears_in_unique_values(self, client, proj):
        """편집으로 생긴 새 값이 드로어 Select 드롭다운에 나와야 한다."""
        _seed(client)
        r = client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "세종"})
        col = next(c for c in r.json()["data"]["meta"]["columns"] if c["key"] == "지역")
        assert "세종" in col["unique_values"]

    def test_prev_records_pipeline_value_not_previous_edit(self, client, proj):
        """같은 칸을 두 번 고쳐도 prev 는 **파이프라인 산출값**이어야 한다.

        되돌리기의 목적지가 '직전 편집값'이면 원래 값으로 못 돌아간다.
        """
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산시"})
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        (e,) = _data(client).json()["edits"]
        assert e["value"] == "부산광역시"
        assert e["prev"] == "부산"                  # '부산시' 가 아니다

    def test_unknown_column_rejected(self, client, proj):
        _seed(client)
        r = client.patch(f"/api/projects/{NAME}/rows/r3", json={"없는컬럼": "x"})
        assert r.status_code == 400
        assert "알 수 없는 컬럼" in r.json()["detail"]

    def test_row_id_not_editable(self, client, proj):
        _seed(client)
        r = client.patch(f"/api/projects/{NAME}/rows/r3", json={ROW_ID_COL: "r99"})
        assert r.status_code == 400

    def test_missing_row_404(self, client, proj):
        _seed(client)
        r = client.patch(f"/api/projects/{NAME}/rows/r999", json={"지역": "x"})
        assert r.status_code == 404

    def test_empty_payload_400(self, client, proj):
        _seed(client)
        assert client.patch(f"/api/projects/{NAME}/rows/r3", json={}).status_code == 400

    def test_before_export_400(self, client, proj):
        """정제·내보내기 전에는 편집할 대상이 없다."""
        r = client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "x"})
        assert r.status_code == 400
        assert "먼저" in r.json()["detail"]

    def test_path_traversal_rejected(self, client, proj):
        r = client.patch("/api/projects/..%2F..%2Fetc/rows/r3", json={"지역": "x"})
        assert r.status_code in (400, 404)


# ─────────────────────────────────────────────────────────────────────────────
# 뒤처짐 판정 (freshness)
# ─────────────────────────────────────────────────────────────────────────────

class TestFreshness:
    def test_fresh_after_rebuild(self, client, proj):
        _seed(client)
        f = client.get(f"/api/projects/{NAME}/freshness").json()
        assert f["is_stale"] is False
        assert f["edit_count"] == 0

    def test_edit_makes_xlsx_stale(self, client, proj):
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        f = client.get(f"/api/projects/{NAME}/freshness").json()
        assert f["is_stale"] is True                # xlsx 가 뒤처졌다
        assert f["edit_count"] == 1
        assert f["overrides_updated_at"] is not None

    def test_rebuild_clears_stale(self, client, proj):
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        client.post(f"/api/projects/{NAME}/rebuild")
        f = client.get(f"/api/projects/{NAME}/freshness").json()
        assert f["is_stale"] is False
        assert f["edit_count"] == 1                 # 편집은 그대로 남아 있다


# ─────────────────────────────────────────────────────────────────────────────
# rebuild / download
# ─────────────────────────────────────────────────────────────────────────────

class TestRebuildAndDownload:
    def test_rebuild_bakes_edit_into_xlsx(self, client, proj):
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        before = _xlsx_mtime(proj)
        r = client.post(f"/api/projects/{NAME}/rebuild")
        assert r.status_code == 200
        assert r.json()["conflicts"] == []
        assert _xlsx_mtime(proj) != before

        ws = openpyxl.load_workbook(proj / "out" / "r.xlsx")["Cleaned"]
        vals = [row[1].value for row in list(ws.rows)[2:]]
        assert vals == ["서울", "부산광역시", "서울", "부산", "서울", "대구"]

    def test_rebuild_reports_conflicts(self, client, proj):
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        # 원본에 없는 행을 가리키도록 overrides 를 직접 조작
        from engine.overrides import Edit as _E, Overrides as _O, save_overrides as _s
        _s(proj, _O(edits=[_E(row_id="r99", col="지역", value="제주")]))
        r = client.post(f"/api/projects/{NAME}/rebuild")
        (c,) = r.json()["conflicts"]
        assert c["reason"] == "row_missing"

    def test_download_rebuilds_when_stale(self, client, proj):
        """★ 다운로드한 엑셀에 편집이 들어 있어야 한다 (요구사항 2번 정방향)."""
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        assert client.get(f"/api/projects/{NAME}/freshness").json()["is_stale"] is True

        r = client.get(f"/api/projects/{NAME}/download")
        assert r.status_code == 200
        # 다운로드가 rebuild 를 유발해 더 이상 뒤처지지 않는다
        assert client.get(f"/api/projects/{NAME}/freshness").json()["is_stale"] is False

        ws = openpyxl.load_workbook(proj / "out" / "r.xlsx")["Cleaned"]
        assert [row[1].value for row in list(ws.rows)[2:]][1] == "부산광역시"


# ─────────────────────────────────────────────────────────────────────────────
# 되돌리기
# ─────────────────────────────────────────────────────────────────────────────

class TestRevert:
    def test_revert_single_cell_restores_pipeline_value(self, client, proj):
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        r = client.request("DELETE", f"/api/projects/{NAME}/overrides",
                           json={"row_id": "r3", "col": "지역"})
        assert r.status_code == 200
        assert r.json()["edit_count"] == 0

        data = json.loads((proj / f"{NAME}_data.json").read_text(encoding="utf-8"))
        assert next(x["지역"] for x in data["rows"] if x[ROW_ID_COL] == "r3") == "부산"

    def test_revert_whole_row(self, client, proj):
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r2", json={"이름": "홍길순", "지역": "인천"})
        r = client.request("DELETE", f"/api/projects/{NAME}/overrides", json={"row_id": "r2"})
        assert r.json()["removed"] == 2
        assert r.json()["edit_count"] == 0

    def test_revert_all(self, client, proj):
        _seed(client)
        client.patch(f"/api/projects/{NAME}/rows/r2", json={"지역": "인천"})
        client.patch(f"/api/projects/{NAME}/rows/r3", json={"지역": "부산광역시"})
        r = client.request("DELETE", f"/api/projects/{NAME}/overrides")
        assert r.json()["removed"] == 2
        assert client.get(f"/api/projects/{NAME}/overrides").json()["count"] == 0


# ─────────────────────────────────────────────────────────────────────────────
# 역방향 xlsx 업로드 (§5.4)
# ─────────────────────────────────────────────────────────────────────────────

class TestImportXlsx:
    def _download_and_edit(self, proj, mutate):
        """현재 cleaned.xlsx 를 읽어 mutate 로 고친 새 파일을 만든다 (사용자가 엑셀로 하는 일)."""
        src = proj / "out" / "r.xlsx"
        wb = openpyxl.load_workbook(src)
        mutate(wb["Cleaned"])
        out = proj / "edited.xlsx"
        wb.save(out)
        return out

    def _post(self, client, path, apply=False):
        with open(path, "rb") as f:
            return client.post(
                f"/api/projects/{NAME}/import-xlsx?apply={str(apply).lower()}",
                files={"file": ("edited.xlsx", f,
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )

    def test_preview_does_not_save(self, client, proj):
        """★ 확인 전에는 아무것도 반영하지 않는다 — 엑셀은 값을 조용히 망가뜨린다."""
        _seed(client)
        edited = self._download_and_edit(proj, lambda ws: ws.cell(4, 2, "부산광역시"))  # r3 의 지역

        r = self._post(client, edited, apply=False)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["status"] == "preview"
        assert body["count"] == 1
        assert body["changes"][0]["col"] == "지역"
        assert body["changes"][0]["value"] == "부산광역시"
        # 저장은 안 됐다
        assert client.get(f"/api/projects/{NAME}/overrides").json()["count"] == 0

    def test_apply_absorbs_changes(self, client, proj):
        _seed(client)
        edited = self._download_and_edit(proj, lambda ws: ws.cell(4, 2, "부산광역시"))

        r = self._post(client, edited, apply=True)
        assert r.status_code == 200, r.text
        assert r.json()["edit_count"] == 1
        (e,) = client.get(f"/api/projects/{NAME}/overrides").json()["edits"]
        assert (e["col"], e["value"], e["prev"], e["origin"]) == (
            "지역", "부산광역시", "부산", "xlsx-import",
        )

    def test_absorbed_edit_survives_rebuild(self, client, proj):
        """엑셀에서 가져온 편집도 정제 재실행을 견뎌야 한다 (드로어 편집과 같은 경로)."""
        _seed(client)
        edited = self._download_and_edit(proj, lambda ws: ws.cell(4, 2, "부산광역시"))
        self._post(client, edited, apply=True)
        client.post(f"/api/projects/{NAME}/rebuild")

        ws = openpyxl.load_workbook(proj / "out" / "r.xlsx")["Cleaned"]
        assert [r[1].value for r in list(ws.rows)[2:]][1] == "부산광역시"

    def test_no_change_returns_empty(self, client, proj):
        _seed(client)
        untouched = self._download_and_edit(proj, lambda ws: None)
        assert self._post(client, untouched).json()["count"] == 0

    def test_deleted_row_id_column_rejected(self, client, proj):
        """★ 숨은 __row_id 열을 지우면 어느 행인지 알 수 없다 — 거부한다."""
        _seed(client)

        def drop_rid(ws):
            ws.delete_cols(4)  # 이름·지역·기관코드·__row_id 중 마지막

        broken = self._download_and_edit(proj, drop_rid)
        r = self._post(client, broken)
        assert r.status_code == 400
        assert ROW_ID_COL in r.json()["detail"]

    def test_unknown_row_id_rejected(self, client, proj):
        _seed(client)
        broken = self._download_and_edit(proj, lambda ws: ws.cell(3, 4, "r999"))
        r = self._post(client, broken)
        assert r.status_code == 400
        assert "r999" in r.json()["detail"]

    def test_duplicated_row_rejected(self, client, proj):
        """행을 복사해 붙여넣은 경우 — 행 추가는 아직 지원하지 않는다."""
        _seed(client)
        broken = self._download_and_edit(proj, lambda ws: ws.cell(4, 4, "r2"))
        r = self._post(client, broken)
        assert r.status_code == 400
        assert "중복" in r.json()["detail"]
