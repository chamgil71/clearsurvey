"""
JSON 업로드 지원 테스트 (배열-of-객체만 지원).

Coverage targets:
  - JsonAnalyzer / json_records_to_columns / load_json_records
  - engine.pipeline._json_to_dataframe
  - SurveyPipeline.run() with json 입력 — xlsx 입력과 결과가 동일해야 한다

실행: pytest tests/test_config_json.py -v
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from engine.config import SurveyConfig
from engine.json_analyzer import (
    JsonAnalyzer,
    JsonFormatError,
    json_records_to_columns,
    load_json_records,
)
from engine.pipeline import SurveyPipeline


_RECORDS = [
    {"이름": "홍길동", "부서": "개발팀", "점수": 90, "연락처": "01012345678"},
    {"이름": "김철수", "부서": "영업팀", "점수": 85, "연락처": "01098765432"},
    {"이름": "이영희", "부서": "개발팀", "점수": 92, "연락처": "01055512345"},
]


def _write_json(path: Path, records=None) -> Path:
    path.write_text(json.dumps(records if records is not None else _RECORDS, ensure_ascii=False), encoding="utf-8")
    return path


def _cfg_for(data_path: Path, tmp_path: Path) -> SurveyConfig:
    return SurveyConfig(
        project="json_test",
        source={"sheet": None, "header_row": 1, "data_start_row": 2, "file": str(data_path)},
        paths={"output_dir": str(tmp_path / "output"), "output_file": "result.xlsx"},
        sheets={"cleaned": "Cleaned", "summary": "Summary"},
        columns=[
            {"output_col": "이름", "source_col": 1, "transform": "copy"},
            {"output_col": "부서", "source_col": 2, "transform": "copy"},
            {"output_col": "점수", "source_col": 3, "transform": "copy"},
            {"output_col": "연락처", "source_col": 4, "transform": "copy"},
        ],
    )


# ─────────────────────────────────────────────────────────────────────────────
# load_json_records / json_records_to_columns
# ─────────────────────────────────────────────────────────────────────────────

class TestLoadJsonRecords:
    def test_rejects_top_level_object(self, tmp_path):
        path = tmp_path / "bad.json"
        path.write_text(json.dumps({"이름": "홍길동"}), encoding="utf-8")
        with pytest.raises(JsonFormatError):
            load_json_records(path)

    def test_rejects_array_of_non_objects(self, tmp_path):
        path = tmp_path / "bad2.json"
        path.write_text(json.dumps([1, 2, 3]), encoding="utf-8")
        with pytest.raises(JsonFormatError):
            load_json_records(path)

    def test_accepts_array_of_objects(self, tmp_path):
        path = _write_json(tmp_path / "ok.json")
        records = load_json_records(path)
        assert len(records) == 3

    def test_empty_array_is_valid(self, tmp_path):
        path = _write_json(tmp_path / "empty.json", records=[])
        assert load_json_records(path) == []


class TestJsonRecordsToColumns:
    def test_header_is_first_seen_key_union(self):
        records = [{"a": 1, "b": 2}, {"b": 3, "c": 4}]
        headers, rows = json_records_to_columns(records)
        assert headers == ["a", "b", "c"]

    def test_missing_key_becomes_none(self):
        records = [{"a": 1, "b": 2}, {"b": 3, "c": 4}]
        headers, rows = json_records_to_columns(records)
        # 두 번째 레코드엔 "a"가 없다 → None
        assert rows[1][headers.index("a")] is None
        # 첫 번째 레코드엔 "c"가 없다 → None
        assert rows[0][headers.index("c")] is None


# ─────────────────────────────────────────────────────────────────────────────
# JsonAnalyzer
# ─────────────────────────────────────────────────────────────────────────────

class TestJsonAnalyzer:
    def test_sheet_names_returns_virtual_sheet(self, tmp_path):
        path = _write_json(tmp_path / "data.json")
        assert JsonAnalyzer(path).sheet_names() == ["JSON"]

    def test_detect_header_row(self, tmp_path):
        path = _write_json(tmp_path / "data.json")
        detection = JsonAnalyzer(path).detect_header_row()
        assert detection["header_row"] == 1
        assert detection["data_start_row"] == 2
        assert detection["column_count"] == 4

    def test_all_headers(self, tmp_path):
        path = _write_json(tmp_path / "data.json")
        assert JsonAnalyzer(path).all_headers() == ["이름", "부서", "점수", "연락처"]

    def test_generate_config_yaml_suggests_phone_transform(self, tmp_path):
        import yaml
        path = _write_json(tmp_path / "data.json")
        analyzer = JsonAnalyzer(path)
        out = analyzer.generate_config_yaml(tmp_path / "config.yaml", project_name="p")
        cfg = yaml.safe_load(out.read_text(encoding="utf-8"))
        phone_col = next(c for c in cfg["columns"] if c["output_col"] == "연락처")
        assert phone_col["transform"] == "normalize_phone"


# ─────────────────────────────────────────────────────────────────────────────
# SurveyPipeline.run() with json — xlsx 결과와 동등해야 한다
# ─────────────────────────────────────────────────────────────────────────────

class TestPipelineJsonVsXlsxEquivalence:
    def test_json_pipeline_produces_same_cleaned_rows_as_xlsx(self, tmp_path):
        xlsx_path = tmp_path / "data.xlsx"
        wb = Workbook()
        ws = wb.active
        headers = ["이름", "부서", "점수", "연락처"]
        for ci, h in enumerate(headers, 1):
            ws.cell(1, ci, h)
        for ri, rec in enumerate(_RECORDS, 2):
            ws.cell(ri, 1, rec["이름"])
            ws.cell(ri, 2, rec["부서"])
            ws.cell(ri, 3, rec["점수"])
            ws.cell(ri, 4, rec["연락처"])
        wb.save(xlsx_path)

        json_path = _write_json(tmp_path / "data.json")

        cfg_xlsx = _cfg_for(xlsx_path, tmp_path / "xlsx_out")
        cfg_json = _cfg_for(json_path, tmp_path / "json_out")

        out_xlsx = SurveyPipeline(cfg_xlsx).run(input_path=xlsx_path)
        out_json = SurveyPipeline(cfg_json).run(input_path=json_path)

        def _cleaned_rows(path: Path) -> list[tuple]:
            wb = openpyxl.load_workbook(path)
            ws = wb["Cleaned"]
            return [
                r for r in ws.iter_rows(min_row=3, max_col=4, values_only=True)
                if any(c is not None for c in r)
            ]

        assert _cleaned_rows(out_xlsx) == _cleaned_rows(out_json)

    def test_json_phone_field_keeps_leading_zero(self, tmp_path):
        """JSON 숫자 리터럴은 앞자리 0을 허용하지 않으므로 애초에 문자열로 온다 — 손실 위험이 없다."""
        json_path = _write_json(tmp_path / "data.json")
        cfg = _cfg_for(json_path, tmp_path)
        out = SurveyPipeline(cfg).run(input_path=json_path)
        wb = openpyxl.load_workbook(out)
        ws = wb["Cleaned"]
        phones = [
            r[3] for r in ws.iter_rows(min_row=3, max_col=4, values_only=True)
            if r[3] is not None
        ]
        assert "01012345678" in phones

    def test_json_numeric_field_is_numeric_type(self, tmp_path):
        json_path = _write_json(tmp_path / "data.json")
        cfg = _cfg_for(json_path, tmp_path)
        out = SurveyPipeline(cfg).run(input_path=json_path)
        wb = openpyxl.load_workbook(out)
        ws = wb["Cleaned"]
        scores = [
            r[2] for r in ws.iter_rows(min_row=3, max_col=4, values_only=True)
            if r[2] is not None
        ]
        assert all(isinstance(s, (int, float)) for s in scores)
        assert 90 in scores

    def test_json_missing_key_produces_empty_cell(self, tmp_path):
        """레코드마다 키가 다를 때 없는 키는 빈 값으로 채워져야 한다(에러 없이)."""
        records = [
            {"이름": "홍길동", "부서": "개발팀", "점수": 90, "연락처": "01012345678"},
            {"이름": "김철수", "점수": 85},  # 부서·연락처 누락
        ]
        json_path = _write_json(tmp_path / "data.json", records=records)
        cfg = _cfg_for(json_path, tmp_path)
        out = SurveyPipeline(cfg).run(input_path=json_path)
        wb = openpyxl.load_workbook(out)
        ws = wb["Cleaned"]
        rows = [
            r for r in ws.iter_rows(min_row=3, max_col=4, values_only=True)
            if any(c is not None for c in r)
        ]
        assert len(rows) == 2
        assert rows[1][0] == "김철수"
        assert rows[1][1] is None  # 부서 누락
