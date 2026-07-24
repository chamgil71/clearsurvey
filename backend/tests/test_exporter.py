"""
Tests for engine/exporter.py

Coverage targets:
  - _detect_type()             — 값 목록 → 타입 자동 감지
  - _build_default_dashboard() — 컬럼 메타 → DashboardConfig 생성
  - export_to_json()           — xlsx → JSON dict 변환
  - _update_manifest()         — projects.json upsert

실행: pytest tests/test_exporter.py -v
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from engine.config import SurveyConfig
from engine.exporter import (
    _build_default_dashboard,
    _detect_type,
    _update_manifest,
    export_to_json,
)


# ─────────────────────────────────────────────────────────────────────────────
# _detect_type — 값 목록 → "numeric" | "category" | "text"
# ─────────────────────────────────────────────────────────────────────────────

class TestDetectType:
    def test_empty_list_returns_text(self):
        assert _detect_type([]) == "text"

    def test_all_none_returns_text(self):
        assert _detect_type([None, None, None]) == "text"

    def test_all_integers_is_numeric(self):
        assert _detect_type([1, 2, 3, 4, 5]) == "numeric"

    def test_all_floats_is_numeric(self):
        assert _detect_type([1.0, 2.5, 3.7]) == "numeric"

    def test_mixed_numeric_and_none_is_numeric(self):
        # 80% 이상 숫자 → numeric
        assert _detect_type([1, None, 3, None, 5]) == "numeric"

    def test_low_unique_ratio_is_category(self):
        # 10개 값, 고유값 2개 → ratio=0.2 ≤ 0.5 → category
        vals = ["A", "B"] * 5
        assert _detect_type(vals) == "category"

    def test_high_unique_ratio_is_text(self):
        # 5개 값, 모두 고유 → ratio=1.0 > 0.5 → text
        assert _detect_type(["apple", "banana", "cherry", "durian", "elderberry"]) == "text"

    def test_too_many_unique_values_is_text(self):
        # 41개 고유값 → unique > 40 → text
        vals = [str(i) for i in range(41)]
        assert _detect_type(vals) == "text"

    def test_boundary_unique_count_40_is_category(self):
        # 고유값 정확히 40개, 충분한 반복 → category
        # 각 값 3번 반복해 ratio ≈ 0.33
        vals = [str(i) for i in range(40)] * 3
        assert _detect_type(vals) == "category"


# ─────────────────────────────────────────────────────────────────────────────
# _detect_type — 리커트/평점 척도(1~5 등) 오탐 방지
# ─────────────────────────────────────────────────────────────────────────────

class TestDetectTypeRatingScale:
    def test_likert_1_to_5_is_category_not_numeric(self):
        # 만족도 1~5점, 반복 응답 → "OO 합계" KPI가 자동 생성되면 안 됨
        vals = [1, 2, 3, 4, 5, 3, 2, 1, 4, 5, 3, 2] * 3
        assert _detect_type(vals) == "category"

    def test_nps_0_to_10_is_category(self):
        vals = ([0, 5, 10, 7, 8, 9, 6, 5, 4, 3, 2, 1] * 3)
        assert _detect_type(vals) == "category"

    def test_likert_scale_string_values_also_category(self):
        # 리커트 문항이 문자열로 저장돼도(예: 엑셀 텍스트 서식) 기존 low_unique_ratio 규칙으로
        # 이미 category였다 — 이번 변경으로 numeric 경로가 하나 더 생겼을 뿐, 기존 경로는 그대로.
        vals = ["1", "2", "3", "4", "5"] * 4
        assert _detect_type(vals) == "category"

    def test_binary_0_1_flag_stays_numeric(self):
        # O_ 접두사 이진 컬럼(0/1) — multibar 자동 구성이 num_cols를 전제하므로 유지돼야 한다.
        vals = [0, 1, 1, 0, 1, 0, 0, 1, 1, 1]
        assert _detect_type(vals) == "numeric"

    def test_sequential_scores_without_repetition_stay_numeric(self):
        # 반복 없이 각 값이 한 번씩만 나오면(척도라기보다 개별 실측값) numeric 유지
        assert _detect_type([1, 2, 3, 4, 5]) == "numeric"

    def test_score_out_of_100_stays_numeric(self):
        # 상한이 10을 넘으면(리커트 범위 밖) numeric 유지
        vals = [78, 85, 90, 92, 78, 85, 90, 92, 78, 85]
        assert _detect_type(vals) == "numeric"

    def test_year_column_stays_numeric(self):
        # 낮은 자리에서 시작하지 않는 컬럼(연도)은 리커트로 오인하면 안 된다.
        vals = [2020, 2021, 2022, 2023, 2020, 2021, 2022, 2023, 2020, 2021]
        assert _detect_type(vals) == "numeric"

    def test_average_score_with_decimals_stays_numeric(self):
        # 소수가 섞이면 평균 점수 같은 실측값일 가능성이 높아 category로 바꾸지 않는다.
        vals = [1.5, 2.0, 3.5, 4.0, 1.5, 2.0, 3.5, 4.0, 1.5, 2.0]
        assert _detect_type(vals) == "numeric"

    def test_too_many_unique_values_for_rating_scale_stays_numeric(self):
        # 고유값이 11개를 넘으면(리커트/NPS 범위 밖) numeric 유지
        vals = list(range(0, 15)) * 2
        assert _detect_type(vals) == "numeric"


# ─────────────────────────────────────────────────────────────────────────────
# _build_default_dashboard — 컬럼 목록 → DashboardConfig 구조
# ─────────────────────────────────────────────────────────────────────────────

def _make_cols(types: list[str], unique_count: int = 3) -> list[dict]:
    cols = []
    for i, t in enumerate(types):
        col: dict = {"key": f"col{i}", "label": f"컬럼{i}", "type": t}
        if t == "category":
            col["unique_count"] = unique_count
            col["unique_values"] = [f"v{j}" for j in range(unique_count)]
        elif t == "numeric":
            col["min"] = 0.0
            col["max"] = 100.0
            col["sum"] = 500.0
        cols.append(col)
    return cols


class TestBuildDefaultDashboard:
    def test_top_level_keys_present(self):
        dash = _build_default_dashboard([])
        assert set(dash.keys()) >= {"version", "kpi", "charts", "list"}

    def test_version_is_1(self):
        assert _build_default_dashboard([])["version"] == 1

    def test_total_rows_kpi_always_first(self):
        dash = _build_default_dashboard([])
        assert dash["kpi"][0]["type"] == "total_rows"

    def test_numeric_col_adds_sum_kpi(self):
        cols = _make_cols(["numeric"])
        cols[0]["key"] = "score"
        cols[0]["label"] = "점수"
        dash = _build_default_dashboard(cols)
        sum_kpi = next((k for k in dash["kpi"] if k["type"] == "sum"), None)
        assert sum_kpi is not None
        assert sum_kpi["col"] == "score"

    def test_category_col_appears_in_charts(self):
        cols = _make_cols(["category"])
        dash = _build_default_dashboard(cols)
        assert any(c.get("col") == "col0" for c in dash["charts"])

    def test_donut_for_up_to_5_unique(self):
        cols = _make_cols(["category"], unique_count=5)
        dash = _build_default_dashboard(cols)
        chart = next(c for c in dash["charts"] if c.get("col") == "col0")
        assert chart["type"] == "donut"

    def test_bar_for_6_to_15_unique(self):
        cols = _make_cols(["category"], unique_count=10)
        dash = _build_default_dashboard(cols)
        chart = next(c for c in dash["charts"] if c.get("col") == "col0")
        assert chart["type"] == "bar"

    def test_hbar_for_over_15_unique(self):
        cols = _make_cols(["category"], unique_count=20)
        dash = _build_default_dashboard(cols)
        chart = next(c for c in dash["charts"] if c.get("col") == "col0")
        assert chart["type"] == "hbar"

    def test_kpi_capped_at_four(self):
        # 총응답수 + 수치 컬럼 5개 → 최대 4개
        cols = _make_cols(["numeric"] * 5)
        dash = _build_default_dashboard(cols)
        assert len(dash["kpi"]) <= 4

    def test_charts_capped_at_six(self):
        cols = _make_cols(["category"] * 10, unique_count=3)
        dash = _build_default_dashboard(cols)
        assert len(dash["charts"]) <= 6

    def test_list_visible_cols_capped_at_eight(self):
        cols = _make_cols(["category"] * 12, unique_count=3)
        dash = _build_default_dashboard(cols)
        assert len(dash["list"]["visible_cols"]) <= 8

    def test_list_filter_cols_capped_at_three(self):
        cols = _make_cols(["category"] * 5, unique_count=3)
        dash = _build_default_dashboard(cols)
        assert len(dash["list"]["filter_cols"]) <= 3

    def test_o_prefix_binary_cols_produce_multibar(self):
        cols = [
            {"key": "O_A", "label": "O_A", "type": "numeric", "min": 0.0, "max": 1.0, "sum": 5.0},
            {"key": "O_B", "label": "O_B", "type": "numeric", "min": 0.0, "max": 1.0, "sum": 3.0},
        ]
        dash = _build_default_dashboard(cols)
        assert any(c.get("type") == "multibar" for c in dash["charts"])

    def test_single_o_col_does_not_produce_multibar(self):
        # O_ 컬럼이 1개뿐이면 multibar 안 만들어짐
        cols = [{"key": "O_A", "label": "O_A", "type": "numeric", "min": 0.0, "max": 1.0, "sum": 3.0}]
        dash = _build_default_dashboard(cols)
        assert not any(c.get("type") == "multibar" for c in dash["charts"])


# ─────────────────────────────────────────────────────────────────────────────
# export_to_json — xlsx Cleaned 시트 → JSON dict
# ─────────────────────────────────────────────────────────────────────────────

@pytest.fixture
def cleaned_xlsx(tmp_path: Path) -> Path:
    """Row 1: section header (skip), Row 2: headers, Row 3+: data."""
    path = tmp_path / "result.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cleaned"
    ws.cell(1, 1, "섹션헤더")   # Row 1: 섹션 행 (내보내기 시 스킵)
    ws.cell(2, 1, "이름")        # Row 2: 출력 헤더
    ws.cell(2, 2, "부서")
    ws.cell(2, 3, "점수")
    data = [
        ("홍길동", "개발팀", 90),
        ("김철수", "영업팀", 85),
        ("이영희", "개발팀", 92),
        ("박민준", "영업팀", 78),
    ]
    for ri, row in enumerate(data, 3):
        for ci, val in enumerate(row, 1):
            ws.cell(ri, ci, val)
    wb.save(path)
    return path


@pytest.fixture
def minimal_cfg() -> SurveyConfig:
    return SurveyConfig(
        project="test_proj",
        sheets={"cleaned": "Cleaned", "summary": "Summary"},
        columns=[],
        paths={"output_dir": "output", "output_file": "result.xlsx"},
    )


class TestExportToJson:
    def test_returns_dict_with_required_keys(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        for key in ("meta", "rows", "aggregates", "dashboard"):
            assert key in result, f"'{key}' missing from result"

    def test_meta_project_name(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        assert result["meta"]["project"] == "test_proj"

    def test_meta_total_rows(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        assert result["meta"]["total_rows"] == 4

    def test_column_metadata_contains_all_headers(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        keys = {c["key"] for c in result["meta"]["columns"]}
        assert {"이름", "부서", "점수"}.issubset(keys)

    def test_score_column_detected_as_numeric(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        score = next(c for c in result["meta"]["columns"] if c["key"] == "점수")
        assert score["type"] == "numeric"
        assert "min" in score
        assert "max" in score
        assert "sum" in score

    def test_dept_column_detected_as_category(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        dept = next(c for c in result["meta"]["columns"] if c["key"] == "부서")
        assert dept["type"] == "category"
        assert "unique_count" in dept
        assert "unique_values" in dept

    def test_aggregates_category_counts(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        agg = result["aggregates"].get("부서", {})
        assert agg.get("개발팀") == 2
        assert agg.get("영업팀") == 2

    def test_numeric_column_not_in_aggregates(self, tmp_path, minimal_cfg):
        path = tmp_path / "result_many_vals.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cleaned"
        ws.cell(1, 1, "섹션")
        ws.cell(2, 1, "점수")
        for i in range(1, 35):
            ws.cell(i + 2, 1, i)
        wb.save(path)
        result = export_to_json(path, minimal_cfg)
        assert "점수" not in result["aggregates"]

    def test_likert_column_gets_donut_not_sum_kpi_end_to_end(self, tmp_path, minimal_cfg):
        """만족도 1~5점(정수 저장) 컬럼이 자동 대시보드에서 '합계' KPI가 아니라 카운트 차트로
        나와야 한다 — 이 테스트가 실패하면 "만족도 합계 327" 같은 KPI가 되돌아온 것이다."""
        path = tmp_path / "result_likert.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cleaned"
        ws.cell(1, 1, "섹션")
        ws.cell(2, 1, "만족도")
        responses = [1, 2, 3, 4, 5, 3, 2, 1, 4, 5, 3, 2, 4, 5, 1]
        for i, v in enumerate(responses, 3):
            ws.cell(i, 1, v)
        wb.save(path)

        result = export_to_json(path, minimal_cfg)
        col = next(c for c in result["meta"]["columns"] if c["key"] == "만족도")
        assert col["type"] == "category"

        dash = result["dashboard"]
        assert not any(k["type"] == "sum" and k.get("col") == "만족도" for k in dash["kpi"])
        chart = next((c for c in dash["charts"] if c.get("col") == "만족도"), None)
        assert chart is not None
        assert chart["type"] == "donut"

        # aggregates도 "1점 N명, 2점 M명…" 형태로 정확히 카운트돼야 한다.
        agg = result["aggregates"]["만족도"]
        assert agg["1"] == 3
        assert agg["5"] == 3

    def test_rows_are_dicts_with_header_keys(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        for row in result["rows"]:
            assert "이름" in row
            assert "부서" in row
            assert "점수" in row

    def test_dashboard_auto_generated_has_correct_structure(self, cleaned_xlsx, minimal_cfg):
        result = export_to_json(cleaned_xlsx, minimal_cfg)
        dash = result["dashboard"]
        assert dash["version"] == 1
        assert isinstance(dash["kpi"], list)
        assert isinstance(dash["charts"], list)
        assert "list" in dash

    def test_writes_json_to_output_path(self, cleaned_xlsx, minimal_cfg, tmp_path):
        out_path = tmp_path / "data" / "test_proj_data.json"
        export_to_json(cleaned_xlsx, minimal_cfg, output_path=out_path)
        assert out_path.exists()
        loaded = json.loads(out_path.read_text(encoding="utf-8"))
        assert loaded["meta"]["project"] == "test_proj"
        assert loaded["meta"]["total_rows"] == 4

    def test_dashboard_json_overrides_auto_generated(
        self, cleaned_xlsx, minimal_cfg, tmp_path
    ):
        """project_dir/dashboard.json 이 있으면 자동 생성 대신 사용된다."""
        proj_dir = tmp_path / "proj"
        proj_dir.mkdir()
        custom = {
            "version": 99, "kpi": [], "charts": [],
            "list": {"visible_cols": [], "filter_cols": []}
        }
        (proj_dir / "dashboard.json").write_text(
            json.dumps(custom), encoding="utf-8"
        )
        result = export_to_json(cleaned_xlsx, minimal_cfg, project_dir=proj_dir)
        assert result["dashboard"]["version"] == 99

    def test_missing_sheet_raises_valueerror(self, tmp_path, minimal_cfg):
        path = tmp_path / "bad.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "WrongSheet"
        wb.save(path)
        with pytest.raises(ValueError, match="Cleaned"):
            export_to_json(path, minimal_cfg)


# ─────────────────────────────────────────────────────────────────────────────
# _update_manifest — projects.json upsert
# ─────────────────────────────────────────────────────────────────────────────

class TestUpdateManifest:
    def test_creates_manifest_when_missing(self, tmp_path, minimal_cfg):
        data_json = tmp_path / "test_proj_data.json"
        data_json.touch()
        _update_manifest(data_json, minimal_cfg)
        manifest = tmp_path / "projects.json"
        assert manifest.exists()
        projects = json.loads(manifest.read_text(encoding="utf-8"))
        ids = {p["id"] for p in projects}
        assert "test_proj" in ids

    def test_upserts_existing_entry(self, tmp_path, minimal_cfg):
        manifest = tmp_path / "projects.json"
        existing = [{"id": "test_proj", "name": "test_proj", "file": "old.json", "updated": "2025-01-01 00:00"}]
        manifest.write_text(json.dumps(existing), encoding="utf-8")

        data_json = tmp_path / "test_proj_data.json"
        data_json.touch()
        _update_manifest(data_json, minimal_cfg)

        projects = json.loads(manifest.read_text(encoding="utf-8"))
        matches = [p for p in projects if p["id"] == "test_proj"]
        assert len(matches) == 1           # 중복 없음
        assert matches[0]["file"] == "test_proj_data.json"

    def test_other_projects_preserved(self, tmp_path, minimal_cfg):
        manifest = tmp_path / "projects.json"
        existing = [{"id": "other_proj", "name": "other_proj", "file": "other.json", "updated": "2025-01-01 00:00"}]
        manifest.write_text(json.dumps(existing), encoding="utf-8")

        data_json = tmp_path / "test_proj_data.json"
        data_json.touch()
        _update_manifest(data_json, minimal_cfg)

        projects = json.loads(manifest.read_text(encoding="utf-8"))
        ids = {p["id"] for p in projects}
        assert "other_proj" in ids        # 기존 프로젝트 보존
        assert "test_proj" in ids
