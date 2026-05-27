from __future__ import annotations

from openpyxl import Workbook, load_workbook

from engine.config import ColumnDef, SurveyConfig
from engine.config_excel import read_config_from_excel, write_config_sheet


def test_config_excel_roundtrip_exclude_preserved(tmp_path):
    """구분=exclude 컬럼은 result.columns에 transform='exclude'로 보존된다."""
    cfg = SurveyConfig(
        columns=[
            ColumnDef(output_col="ID", source_col=1, transform="copy"),
            ColumnDef(output_col="메모", source_col=2, transform="copy"),
        ]
    )
    path = tmp_path / "config.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    write_config_sheet(wb, cfg)
    wb.save(path)

    # 10열 구조: col3=output_col, col2=구분
    wb2 = load_workbook(path)
    ws = wb2["Config"]
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 3).value == "메모":
            ws.cell(row, 2).value = "exclude"
            break
    wb2.save(path)

    result = read_config_from_excel(path)

    output_cols = [col.output_col for col in result.columns]
    assert "ID" in output_cols
    assert "메모" in output_cols
    memo_col = next(c for c in result.columns if c.output_col == "메모")
    assert memo_col.transform == "exclude"


def test_config_excel_roundtrip_source_cols(tmp_path):
    """source_cols는 Excel 라운드트립에서 복원된다. year_col은 YAML 전용."""
    cfg = SurveyConfig(
        columns=[
            ColumnDef(output_col="합계", source_cols=[2, 3, 4], transform="group_sum"),
            ColumnDef(output_col="연도", source_col=5, transform="norm_date"),
        ]
    )
    path = tmp_path / "config.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    write_config_sheet(wb, cfg)
    wb.save(path)

    result = read_config_from_excel(path)

    assert result.columns[0].source_cols == [2, 3, 4]
    # year_col은 Config 시트에 없음 (YAML 전용 필드)
    assert result.columns[1].year_col is None
