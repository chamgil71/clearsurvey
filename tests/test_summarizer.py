from __future__ import annotations

import pandas as pd
from openpyxl import Workbook

from engine.config import (
    SurveyConfig, SummaryConfig, SummarySection, SummaryLayout,
    SummaryTotalItem, SummaryKeyword, SummaryColumnItem, ColumnDef
)
from engine.summarizer import SummarySheetWriter, _resolve_positions
from engine.pipeline import SurveyPipeline


def test_resolve_positions_layout():
    """레이아웃 설정에 따라 각 섹션의 start_row와 start_col이 오버랩 없이 자동 계산된다."""
    # 2개 단(column), 각 단은 4열 span, 1열 간격, 행 간격은 2행, 3행부터 시작
    layout = SummaryLayout(
        cols=2,
        start_row=3,
        col_span=4,
        gap_cols=1,
        gap_rows=2
    )
    
    sections = [
        # 단 1: totals (항목 2개 -> 높이 4)
        SummarySection(
            id="sec1",
            title="요약1",
            type="totals",
            layout_col=1,
            items=[
                SummaryTotalItem(label="전체", type="count_all", col_ref="ID"),
                SummaryTotalItem(label="완료", type="countif_exact", col_ref="상태", value="완료"),
            ]
        ),
        # 단 1: unique_count (cleaned_col_vals 에 고유값 3개 제공 -> 높이 5)
        SummarySection(
            id="sec2",
            title="요약2",
            type="unique_count",
            layout_col=1,
            col_ref="부서"
        ),
        # 단 2: unique_count (start_row/start_col 강제 override)
        SummarySection(
            id="sec3",
            title="요약3",
            type="unique_count",
            layout_col=2,
            col_ref="직급",
            start_row=5,
            start_col=10
        ),
        # 단 2: binary_sum (항목 2개 -> 높이 4)
        SummarySection(
            id="sec4",
            title="요약4",
            type="binary_sum",
            layout_col=2,
            columns=[
                SummaryColumnItem(col_ref="만족"),
                SummaryColumnItem(col_ref="추천"),
            ]
        )
    ]
    
    cfg = SurveyConfig(
        project="test_proj",
        columns=[],
        summary=SummaryConfig(layout=layout, sections=sections)
    )
    
    df = pd.DataFrame()
    col_index_map = {"ID": 1, "상태": 2, "부서": 3, "직급": 4, "만족": 5, "추천": 6}
    cleaned_col_vals = {
        "부서": ["영업", "개발", "기획"]  # 3개 고유값
    }
    
    resolved = _resolve_positions(cfg, df, col_index_map, cleaned_col_vals)
    
    # sec1 검증 (단 1의 첫 섹션)
    # start_row = 3, start_col = 1
    assert resolved[0].start_row == 3
    assert resolved[0].start_col == 1
    
    # sec2 검증 (단 1의 두 번째 섹션)
    # sec1 높이: title(1) + header(1) + 2 items = 4
    # sec2 start_row = 3 + 4 + gap_rows(2) = 9
    # sec2 start_col = 1
    assert resolved[1].start_row == 9
    assert resolved[1].start_col == 1
    
    # sec3 검증 (단 2의 첫 섹션이지만 manual override가 있으므로 그대로 유지)
    assert resolved[2].start_row == 5
    assert resolved[2].start_col == 10
    
    # sec4 검증 (단 2의 두 번째 섹션)
    # sec3(override) start_row=5 이고 높이는 unique_count 데이터 없으므로 2
    # cursors[2] = 5 + 2 + gap_rows(2) = 9
    # sec4 start_row = 9
    # sec4 start_col = 1 + (2-1)*(4+1) = 6
    assert resolved[3].start_row == 9
    assert resolved[3].start_col == 6


def test_summarizer_sheet_writer_smoke():
    """SummarySheetWriter가 각 섹션 타입에 맞는 수식 및 데이터를 오류 없이 작성한다."""
    cfg = SurveyConfig(
        project="test_proj",
        sheets={"cleaned": "정제", "summary": "요약"},
        columns=[
            ColumnDef(output_col="ID", source_col=1, transform="copy"),
            ColumnDef(output_col="부서", source_col=2, transform="copy"),
            ColumnDef(output_col="만족도", source_col=3, transform="to_pct"),
            ColumnDef(output_col="추천여부", source_col=4, transform="to_binary"),
        ],
        summary=SummaryConfig(
            sheet_name="요약",
            total_cell="$B$3",
            sections=[
                SummarySection(
                    title="전체 현황",
                    type="totals",
                    start_row=1,
                    start_col=1,
                    items=[
                        SummaryTotalItem(label="전체 건수", type="count_all", col_ref="ID"),
                        SummaryTotalItem(label="만족 건수", type="countif_exact", col_ref="추천여부", value="1"),
                    ]
                ),
                SummarySection(
                    title="부서별 분포",
                    type="unique_count",
                    start_row=1,
                    start_col=5,
                    col_ref="부서",
                    sort=True
                ),
                SummarySection(
                    title="만족도 집계",
                    type="binary_sum",
                    start_row=10,
                    start_col=1,
                    columns=[
                        SummaryColumnItem(col_ref="추천여부", label="추천"),
                    ]
                ),
                SummarySection(
                    title="키워드 요약",
                    type="countif_contains",
                    start_row=10,
                    start_col=5,
                    col_ref="부서",
                    keywords=[
                        SummaryKeyword(label="개발부서", keyword="개발"),
                    ]
                )
            ]
        )
    )
    
    df = pd.DataFrame([
        [1, "개발팀", 90.0, 1],
        [2, "영업팀", 80.0, 0],
        [3, "개발본부", 95.0, 1],
    ])
    
    col_index_map = {"ID": 1, "부서": 2, "만족도": 3, "추천여부": 4}
    cleaned_col_vals = {
        "ID": [1, 2, 3],
        "부서": ["개발팀", "영업팀", "개발본부"],
        "만족도": [90.0, 80.0, 95.0],
        "추천여부": [1, 0, 1],
    }
    
    wb = Workbook()
    # clean_name sheet must exist because _data_ref references it
    wb.create_sheet("정제")
    
    writer = SummarySheetWriter(cfg)
    writer.write(wb, df, col_index_map, cleaned_col_vals)
    
    assert "요약" in wb.sheetnames
    ws = wb["요약"]
    
    # totals 검증
    assert ws.cell(1, 1).value == "전체 현황"
    assert ws.cell(2, 1).value == "항목"
    assert ws.cell(3, 1).value == "전체 건수"
    assert ws.cell(3, 2).value == "=COUNTA('정제'!A3:A10000)"  # COUNTA 수식
    assert ws.cell(4, 1).value == "만족 건수"
    assert ws.cell(4, 2).value == '=COUNTIF(\'정제\'!D3:D10000,"1")'
    
    # unique_count 검증
    assert ws.cell(1, 5).value == "부서별 분포"
    # 고유값: 개발본부, 개발팀, 영업팀 (sorted)
    assert ws.cell(3, 5).value == "개발본부"
    assert ws.cell(4, 5).value == "개발팀"
    assert ws.cell(5, 5).value == "영업팀"
    assert ws.cell(3, 6).value == '=COUNTIF(\'정제\'!B3:B10000,"개발본부")'


def test_dry_run_pipeline(tmp_path):
    """dry_run=True 시 파일 및 JSON 설정이 디스크에 생성되지 않으며 run은 None을 리턴한다."""
    # 더미 데이터 생성
    data_path = tmp_path / "dummy_data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1, "ID")
    ws.cell(1, 2, "이름")
    ws.cell(2, 1, 1)
    ws.cell(2, 2, "홍길동")
    wb.save(data_path)
    
    cfg = SurveyConfig(
        project="dry_run_test",
        source=dict(sheet=None, header_row=1, data_start_row=2, file=str(data_path)),
        paths=dict(output_dir=str(tmp_path / "output"), output_file="cleaned.xlsx"),
        sheets=dict(cleaned="정제", summary="요약"),
        columns=[
            dict(output_col="ID", source_col=1, transform="copy"),
            dict(output_col="이름", source_col=2, transform="mask_name"),
        ]
    )
    
    config_yaml_path = tmp_path / "config.yaml"
    
    # dry_run=True로 실행
    pipeline = SurveyPipeline(cfg, config_path=config_yaml_path)
    out_path = pipeline.run(input_path=data_path, dry_run=True)
    
    # 디스크 저장 검증
    assert out_path is None
    assert not (tmp_path / "output" / "cleaned.xlsx").exists()
    assert not (tmp_path / "dry_run_test_config.json").exists()
